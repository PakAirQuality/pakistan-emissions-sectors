"""Independent primary-source cells against the final extension exports."""
import collections
import csv
import importlib.util
import json
from pathlib import Path
import re
import subprocess
import sys
import unittest

import openpyxl
from pkinv.assets import asset_path
from pkinv.core.io import read_rows
from pkinv.sectors.transport import ocac
from _support import ROOT
from _support import require_layers  # noqa: E402


class ExtensionInputs(unittest.TestCase):
    def table(self, path):
        return read_rows(asset_path(ROOT, path))

    def test_fiscal_agriculture_full_year_cells_from_later_primary_table(self):
        require_layers("the official extension documents")
        path = Path(ROOT) / "data/raw/finance_division/survey_2025_26/2_Agriculture.pdf"
        text = subprocess.check_output(["pdftotext", "-layout", str(path), "-"]).decode()
        page = next(p for p in text.split("\f") if "PRODUCTION OF IMPORTANT CROPS" in p)
        source = re.search(r"^2024-25\s+(.+)$", page, re.M).group(1).split()
        final = {r["item"]: float(r["value"]) for r in self.table("data/activity/agriculture/es_fiscal_series.csv")
                 if r["fiscal_end_year"] == "2025"}
        for item, index in (("wheat_production_t", 0), ("rice_production_t", 1),
                            ("maize_production_t", 4), ("sugarcane_production_t", 8)):
            self.assertEqual(float(source[index].replace(",", "")) * 1000, final[item])
        rows = self.table("data/activity/agriculture/es_fiscal_series.csv")
        self.assertEqual(110, len(rows))
        self.assertEqual(set(range(2016, 2026)), {int(r["fiscal_end_year"]) for r in rows})
        self.assertEqual(len(rows), len({(r["fiscal_end_year"], r["item"]) for r in rows}))

    def test_final_monthly_qim_matches_independent_cell_coordinates(self):
        require_layers("the official extension documents")
        path = Path(ROOT) / "data/raw/pbs/qim/extension_2026/Table-3-2024-25-6.xlsx"
        book = openpyxl.load_workbook(path, data_only=True, read_only=True)
        sheet = book.active
        rows = [r for r in self.table("data/activity/monthly/pbs_qim_month_item.csv")
                if r["fiscal_year_start"] == "2024"]
        for row_number, serial in ((73, "58"), (96, "82"), (97, "83"), (100, "85")):
            source = next(sheet.iter_rows(min_row=row_number, max_row=row_number, values_only=True))
            members = sorted((r for r in rows if r["item_no"] == serial), key=lambda r: r["month"])
            expected = {d.strftime("%Y%m"): float(q) for d, q in
                        zip(next(sheet.iter_rows(min_row=5, max_row=5, values_only=True))[4:16], source[4:16])}
            self.assertEqual(expected, {r["month"]: float(r["quantity"]) for r in members})
            self.assertAlmostEqual(float(source[16]), sum(float(r["quantity"]) for r in members), delta=2)
        self.assertFalse(any(r["item_no"] == "Manufacturing of Iron & Steel Products" for r in rows))
        book.close()

    def test_new_ocac_header_preserves_both_products_and_final_source_mass(self):
        require_layers("the official extension documents")
        path = Path(ROOT) / "data/raw/ocac/extension_2026/ocac_sector_province_2026_06.pdf"
        page = subprocess.check_output(["pdftotext", "-layout", str(path), "-"]).decode().split("\f")[0]
        self.assertIn("HOBC NAPHTHA", page)
        parsed = ocac.parse_page(page)
        self.assertIn("HOBC", parsed["products"])
        self.assertIn("NAPHTHA", parsed["products"])
        source_line = next(line for line in page.splitlines() if line.strip().startswith("Road "))
        source_cells = source_line.split()[1:]
        hsd = float(source_cells[7].replace(",", ""))
        final = [r for r in self.table("outputs/pakemis_fy2025_26_provisional_activity.csv")
                 if r["month"] == "202606" and r["area"] == "PAKISTAN" and r["activity"] == "Road:HSD"]
        self.assertEqual(1, len(final))
        self.assertEqual(hsd, float(final[0]["quantity"]))

    def test_power_fuel_is_not_zeroed_when_net_generation_is_zero(self):
        rows = self.table("data/activity/power/plant_fuel_ledger.csv")
        matched = [r for r in rows if r["fiscal_year"] == "2024" and r["plant"] == "KAPCO" and r["fuel"] == "HSD"]
        self.assertEqual(1, len(matched))
        self.assertGreater(float(matched[0]["quantity"]), 0)
        self.assertEqual(0, float(matched[0]["net_generation_GWh"]))
        fuel = [r for r in self.table("data/activity/power/fuel_consumed.csv")
                if r["fiscal_year"] == "2024" and r["fuel"] == "Diesel Oil"]
        self.assertGreater(float(fuel[0]["toe"]), 0)

    def test_power_controls_distinguish_unresolved_mass_and_conditional_cases(self):
        for row in self.table("outputs/pakemis_power_extension_controls.csv"):
            if row["estimate_status"] != "derived_fleet_intensity_gap":
                continue
            low, central, high = (float(row[k]) for k in ("gap_energy_low_TJ", "estimated_gap_energy_TJ", "gap_energy_high_TJ"))
            self.assertLessEqual(low - 1e-7, central)
            self.assertLessEqual(central, high + 1e-7)
            self.assertGreater(float(row["unresolved_generation_GWh"]), 0)

    def test_fertilizer_gas_class_matches_full_year_ogra_source(self):
        require_layers("the official extension documents")
        path = Path(ROOT) / "data/raw/ogra/state_of_industry/ogra_petroleum_industry_report_2024_25.pdf"
        text = subprocess.check_output(["pdftotext", "-layout", str(path), "-"]).decode()
        source = re.search(r"^\s*4\s+Fertilizer\s+([\d.]+)\s+([\d.]+)\s+([\d.]+)\s+([\d.]+)\s+([\d.]+)", text, re.M)
        self.assertIsNotNone(source)
        final = [r for r in self.table("data/activity/fugitives/fertilizer_gas_class.csv")
                 if r["fiscal_year"] == "2025"]
        self.assertEqual(1, len(final))
        fields = ("sngpl_mmcfd", "ssgcl_mmcfd", "independent_mmcfd", "total_mmcfd")
        self.assertEqual(list(map(float, source.groups()[:4])), [float(final[0][k]) for k in fields])
        self.assertAlmostEqual(sum(float(final[0][k]) for k in fields[:3]), float(final[0]["total_mmcfd"]))

    def test_aviation_final_panel_closes_own_year_pcaa_control(self):
        require_layers("the official extension documents")
        path = Path(ROOT) / "data/raw/pcaa/extension_2025/major-traffic-flows-by-airports-during-the-year-2024-25.pdf"
        text = subprocess.check_output(["pdftotext", "-layout", str(path), "-"]).decode()
        printed = next(line for line in text.splitlines() if "Grand Totals" in line)
        values = [float(v.replace(",", "")) for v in re.findall(r"[\d,]+(?:\.\d+)?", printed)]
        final = [r for r in self.table("data/activity/aviation/movements_airport.csv") if r["year"] == "2024"]
        self.assertEqual(37, len(final))
        for field, column in (("com_dom_mv", 0), ("com_intl_mv", 1), ("noncom_mv", 3),
                              ("pax_dom", 5), ("pax_intl", 6)):
            self.assertEqual(values[column], sum(float(r[field]) for r in final))
        self.assertEqual(37, len({r["icao"] for r in final}))

    def test_household_new_year_weather_matches_retained_own_year_response(self):
        require_layers("the official extension documents")
        final = [r for r in self.table("data/activity/residential/monthly_degree_days.csv")
                 if "202407" <= r["month"] <= "202506"]
        self.assertEqual(60, len(final))
        self.assertEqual(60, len({(r["province"], r["month"]) for r in final}))
        for row in final:
            path = Path(ROOT) / ("data/raw/nasa_power/T2M_monthly_%s_2024_2025.json" % row["city"])
            source = json.loads(path.read_text())["properties"]["parameter"]["T2M"][row["month"]]
            self.assertEqual(source, float(row["t2m_C"]))
            from calendar import monthrange
            days = monthrange(int(row["month"][:4]), int(row["month"][4:]))[1]
            self.assertEqual(round(max(18 - source, 0) * days, 1), float(row["hdd18"]))
            self.assertEqual(round(max(source - 24, 0) * days, 1), float(row["cdd24"]))

    def test_candidate_estimates_remain_separate_and_close_monthly(self):
        annual = self.table("outputs/pakemis_fy2025_26_provisional_emissions.csv")
        monthly = self.table("outputs/pakemis_fy2025_26_provisional_monthly.csv")
        totals = collections.defaultdict(float)
        periods = collections.defaultdict(set)
        for row in monthly:
            key = row["source_process"], row["species"]
            totals[key] += float(row["emissions_Gg"])
            periods[key].add(row["month"])
        for row in annual:
            key = row["source_process"], row["species"]
            self.assertEqual("FY2025-26", row["period"])
            self.assertIn("provisional", row["status"])
            self.assertEqual(12, len(periods[key]))
            self.assertAlmostEqual(float(row["emissions_Gg"]), totals[key], delta=max(1e-10, float(row["emissions_Gg"]) * 1e-12))
        self.assertFalse(any(r["source_process"].lower() == "national total" for r in annual))

    def test_revision_totals_do_not_silently_replace_unpublished_product_cells(self):
        revised = {r["customer_class"]: r for r in self.table("outputs/pakemis_ocac_extension_revisions.csv")}
        self.assertGreater(float(revised["Road"]["delta_tonnes"]), 0)
        self.assertIn("product_split_unresolved", revised["Road"]["status"])
        core = self.table("data/activity/transport/ocac_sectoral_annual.csv")
        road = sum(float(r["tonnes"]) for r in core if r["fiscal_year_start"] == "2024"
                   and r["sector_row"] == "Road" and r["product"] == "TOTAL")
        self.assertEqual(float(revised["Road"]["old_tonnes"]), road)

    def test_offline_primary_extraction_reproduces_final_exports(self):
        require_layers("the official extension documents")
        result = subprocess.run([sys.executable, "scripts/ingest_release_extension.py", "--check"],
                                cwd=ROOT, text=True, capture_output=True)
        self.assertEqual(0, result.returncode, result.stdout + result.stderr)


if __name__ == "__main__":
    unittest.main()
