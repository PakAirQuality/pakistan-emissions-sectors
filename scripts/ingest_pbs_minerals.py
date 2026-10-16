"""Extract PBS production at its printed year and reporting-area grain.

The compiled workbook named 2020_21 prints 2002-03..2022-23. Its printed-year
blocks are authoritative. Worksheets are selected by schema/provenance, never by
equal numerical values. PDF groups are buffered until Total, allowing centred
mineral labels to name preceding province rows. Unknown cells remain empty.
Run holdings.py minerals before extraction; --check does not write.
"""
from __future__ import annotations
import argparse
import collections
import glob
import os
import re
import subprocess
from pkinv.assets import asset_path
from pkinv.core.io import read_rows, write_or_check
from pkinv.paths import ROOT

RAW = asset_path(ROOT, "data/raw/pbs/minerals")
OUT = asset_path(ROOT, "data/activity/monthly/pbs_minerals_month_province.csv")
QUARANTINE = asset_path(ROOT, "outputs/pbs_minerals_quarantine.csv")
SUMMARY = asset_path(ROOT, "outputs/pbs_minerals_parse_summary.csv")
CONTROLS = asset_path(ROOT, "outputs/pakemis_minerals_controls.csv")
# The reporting window, read rather than typed (rc68): with nine years typed here the parse
# summary said no core year was absent while FY2024-25 had no PBS minerals edition at all.
from pkinv.core.reporting_window import CORE_FY_STARTS  # noqa: E402
WINDOW = tuple(CORE_FY_STARTS)
MONTHS_IN_YEAR, TRAILING_CELLS = 12, 13
ROW_TOLERANCE, GROUP_TOLERANCE = 1.0, 2.0
NATIONAL = "National (edition prints no province split)"
COMPILED = "pbs_minerals_2020_21.xlsx"
SHEETS = {"pbs_minerals_2017_18.xlsx": "17-18", "pbs_minerals_2018_19.xlsx": "Original",
          "pbs_minerals_2019_20.xlsx": "Original", COMPILED: "mineral production"}
PROVINCES = ("Punjab", "Sindh", "KPK", "Khyber Pakhtunkhwa", "Balochistan", "Baluchistan",
             "AJK", "Azad Kashmir", "Gilgit Baltistan", "GB", "Federal", "Islamabad", "FATA")
CANONICAL = {
    "NATURAL GAS": ("natural_gas", "fugitives", "MMCFt", "field production; not deliveries"),
    "GAS": ("natural_gas", "fugitives", "MMCFt", "field production"),
    "CRUDE OIL": ("crude_oil", "fugitives", "US_barrel", "wellhead production"),
    "COAL": ("coal", "industry", "t", "mine production; not combustion timing"),
    "LIME STONE": ("limestone", "ippu", "t", "quarry production; cement cross-check"),
    "LIMESTONE": ("limestone", "ippu", "t", "quarry production; cement cross-check"),
}
FIELDS = ("fiscal_year_start", "month", "mineral", "province", "quantity", "year_total_printed",
          "row_control", "canonical_mineral", "consuming_sector", "temporal_method", "timing_target",
          "basis", "source_file", "quantity_unit", "observation_status", "source_sheet", "source_row",
          "source_header", "publication_status", "group_id", "source_national_month_quantity",
          "source_expected_provinces", "group_month_control", "annual_control", "national_coverage",
          "timing_usable", "quality_note")
BAD_FIELDS = ("fiscal_year_start", "mineral", "province", "months_sum", "year_total_printed",
              "source_file", "source_sheet", "source_row", "reason")
CONTROL_FIELDS = ("fiscal_year_start", "mineral", "canonical_mineral", "source_file", "source_sheet",
                  "source_header", "publication_status", "group_id", "source_expected_provinces",
                  "total_row_printed", "provinces_sum", "max_month_residual", "closed",
                  "national_coverage", "timing_usable", "quality_note", "source_total_months",
                  "selected", "selection_reason")

def canonical_key(name):
    text = re.sub(r"\([^)]*(?:\)|$)", " ", str(name))
    text = re.sub(r"\d+", "", text)
    words = text.replace("-", " ").upper().split()
    if len(words) % 2 == 0 and words[:len(words) // 2] == words[len(words) // 2:]:
        words = words[:len(words) // 2]
    return " ".join(words)

def fiscal_year_of(path):
    return int(re.search(r"_(\d{4})_\d{2}\.(?:pdf|xlsx)$", str(path)).group(1))

def months_for(fy_start):
    return ["%04d%02d" % (fy_start, m) for m in range(7, 13)] + [
        "%04d%02d" % (fy_start + 1, m) for m in range(1, 7)]

def _number(cell):
    if cell is None or str(cell).strip() in ("", "-", "—", "NR", "N.R."):
        return None
    try:
        return float(str(cell).replace(",", "").strip())
    except ValueError:
        return None

def _status(cell):
    if _number(cell) is not None:
        return "observed_zero" if _number(cell) == 0 else "observed"
    value = str(cell or "").strip()
    return ("not_reported" if value.upper() in ("NR", "N.R.") else
            "obscured_in_source" if "#" in value else
            "printed_dash" if value in ("-", "—") else "missing")

def _province(text):
    key = str(text or "").strip()
    if key.lower() == "total":
        return "Total"
    for province in PROVINCES:
        if key.lower() == province.lower():
            return {"Khyber Pakhtunkhwa": "KPK", "Baluchistan": "Balochistan"}.get(province, province)
    return None

def _year_header(text):
    match = re.search(r"MINERALS\s+PRODUCTION\s+DATA\s+(\d{4})\s*-\s*(\d{2})", text, re.I)
    if not match:
        return None
    start = int(match.group(1))
    if (start + 1) % 100 != int(match.group(2)):
        raise ValueError("inconsistent fiscal header: %s" % text)
    return start

def _publication_status(header):
    return "revised" if re.search(r"\(\s*R\s*\)", header, re.I) else (
        "final" if re.search(r"\(\s*F\s*\)", header, re.I) else "not_stated")

def _missing_total(group):
    return {"province": "Total", "values": [None] * 12, "annual": None,
            "source_row": group[-1]["source_row"]}

def _xlsx_groups(path):
    import openpyxl
    book = openpyxl.load_workbook(path, data_only=True, read_only=True)
    name = os.path.basename(path)
    wanted = SHEETS.get(name)
    if wanted not in book.sheetnames:
        raise ValueError("undeclared monthly worksheet in %s; available %s" % (name, book.sheetnames))
    sheet = book[wanted]
    fy, header, province_col, month_cols, total_col = None, "", None, [], None
    group, names, result = [], [], []
    for row_number, values in enumerate(sheet.iter_rows(values_only=True), 1):
        text = " ".join(str(v) for v in values if v is not None)
        printed_year = _year_header(text)
        if printed_year is not None:
            if group:
                result.append((fy, header, wanted, " ".join(names), group, _missing_total(group)))
                group = []
            fy, header = printed_year, text
            province_col, month_cols, total_col, names = None, [], None, []
            continue
        if fy is None:
            continue
        headings = [str(v or "").strip().lower() for v in values]
        if "provinces" in headings:
            province_col = headings.index("provinces")
            month_cols = list(range(province_col + 1, province_col + 13))
            expected = ("jul", "aug", "sep", "oct", "nov", "dec", "jan", "feb", "mar", "apr", "may", "jun")
            if not all(headings[i].startswith(month) for i, month in zip(month_cols, expected)):
                raise ValueError("not a monthly annual table: %s/%s row %d" % (name, wanted, row_number))
            total_col = province_col + 13
            continue
        if province_col is None or fy not in WINDOW:
            continue
        province = _province(values[province_col])
        prefix = " ".join(str(v).strip() for v in values[:province_col]
                          if v is not None and _number(v) is None)
        if prefix:
            if group and names and canonical_key(prefix) != canonical_key(" ".join(names)):
                result.append((fy, header, wanted, " ".join(names), group, _missing_total(group)))
                group, names = [], []
            names.append(prefix)
        if province is None:
            continue
        record = {"province": province, "values": [values[i] for i in month_cols],
                  "annual": values[total_col], "source_row": row_number}
        if province == "Total":
            result.append((fy, header, wanted, " ".join(names), group, record))
            group, names = [], []
        else:
            group.append(record)
    if group:
        result.append((fy, header, wanted, " ".join(names), group, _missing_total(group)))
    book.close()
    return result

def _pdf_groups(path):
    text = subprocess.check_output(["pdftotext", "-layout", path, "-"]).decode("utf-8", "replace")
    lines = text.splitlines()
    header = next((line.strip() for line in lines if _year_header(line) is not None), "")
    fy = _year_header(header)
    if fy is None or fy != fiscal_year_of(path):
        raise ValueError("PDF filename and printed fiscal header disagree: %s" % path)
    width = 12 if fy == 2021 else 13
    province_pattern = re.compile(r"(?<!\w)(%s)(?!\w)" % "|".join(
        sorted(PROVINCES + ("Total",), key=len, reverse=True)), re.I)
    token = re.compile(r"(?:\d[\d,]*(?:\.\d+)?|[-—]|NR|#+)", re.I)
    group, names, result = [], [], []
    pending_province = None
    for line_number, line in enumerate(lines, 1):
        if not line.strip() or _year_header(line) is not None:
            continue
        if any(word in line.lower() for word in ("provinces", "name of", "july,", "grand-")):
            continue
        found = province_pattern.search(line)
        if found:
            prefix, tail = line[:found.start()].strip(), line[found.end():]
            if prefix and re.search(r"[A-Za-z]", prefix):
                if (group and names and canonical_key(prefix)
                        and canonical_key(" ".join(names)) not in ("NATURAL", "CRUDE", "LIME")
                        and canonical_key(prefix) != canonical_key(" ".join(names))):
                    result.append((fy, header, "PDF", " ".join(names), group, _missing_total(group)))
                    group, names = [], []
                names.append(prefix)
            province, row_number = _province(found.group()), line_number
            values = token.findall(tail)
            if not values:
                pending_province = (province, row_number)
                continue
            pending_province = None
        elif pending_province is not None:
            values = token.findall(line)
            if len(values) != width:
                value = re.sub(r"^\s*\d+\s*", "", line).strip()
                if value and re.search(r"[A-Za-z]", value) and value.lower() not in ("s.", "n", "o", "mineral"):
                    names.append(value)
                continue
            province, row_number = pending_province
            pending_province = None
        else:
            value = re.sub(r"^\s*\d+\s*", "", line).strip()
            if value and re.search(r"[A-Za-z]", value) and value.lower() not in ("s.", "n", "o", "mineral"):
                names.append(value)
            continue
        if len(values) != width:
            values = [None] * width  # a short line cannot be shifted into twelve months
        record = {"province": province, "values": values[:12],
                  "annual": values[12] if width == 13 else None, "source_row": row_number}
        if province == "Total":
            result.append((fy, header, "PDF", " ".join(names), group, record))
            group, names, pending_province = [], [], None
        else:
            group.append(record)
    if group:
        result.append((fy, header, "PDF", " ".join(names), group, _missing_total(group)))
    return result

def _assemble(path, raw_group):
    fy, header, sheet, printed_name, provinces, total = raw_group
    mineral = canonical_key(printed_name)
    if not mineral:
        mineral = "UNLABELLED SOURCE GROUP %s" % total["source_row"]
    canonical, sector, unit, basis = CANONICAL.get(mineral, ("", "", "t", ""))
    if not canonical:
        unit = "g" if re.search(r"grams?", printed_name, re.I) else (
            "kg" if re.search(r"\bkg\b", printed_name, re.I) else "t")
    group_id = "%s:%s:%s:%s:%s" % (os.path.basename(path), sheet, fy, mineral, total["source_row"])
    reported_areas = sorted(r["province"] for r in provinces)
    expected_areas = reported_areas
    if canonical:
        expected_areas = ([NATIONAL] if canonical in ("crude_oil", "natural_gas") and fy < 2017
                          else sorted(["Punjab", "Sindh", "KPK", "Balochistan"] +
                                      (["FATA"] if canonical in ("coal", "limestone") and fy < 2018 else [])))
    expected = ";".join(expected_areas) or NATIONAL
    duplicates = len({r["province"] for r in provinces}) != len(provinces)
    national_values = [_number(v) for v in total["values"]]
    complete = all(v is not None for v in national_values) and all(
        _number(v) is not None for r in provinces for v in r["values"])
    residual = (max(abs(sum(_number(r["values"][i]) for r in provinces) - national_values[i])
                    for i in range(12)) if complete and provinces else 0.0 if complete else None)
    annual_ok, bad = True, []
    for record in provinces + [total]:
        vals, printed = [_number(v) for v in record["values"]], _number(record["annual"])
        if all(v is not None for v in vals) and printed is not None and abs(sum(vals) - printed) > ROW_TOLERANCE:
            annual_ok = False
            bad.append(dict(zip(BAD_FIELDS, (fy, mineral, record["province"], sum(vals), printed,
                            os.path.basename(path), sheet, record["source_row"],
                            "monthly sum differs from printed annual control"))))
    actual_areas = reported_areas or [NATIONAL]
    closed = (complete and not duplicates and residual <= GROUP_TOLERANCE and annual_ok
              and actual_areas == expected_areas)
    suspect = fy == 2021 and canonical in ("crude_oil", "natural_gas")
    quality = ("Source repeats anomalous FY2021-22 oil/gas provincial monthly entries; "
               "no independent monthly witness establishes a correction; timing falls back to T4"
               if suspect else "" if closed else "Incomplete, duplicated or non-closing source group")
    timing = closed and not suspect
    admitted = provinces or [dict(total, province=NATIONAL)]
    out = []
    for record in admitted:
        printed = _number(record["annual"])
        for i, month in enumerate(months_for(fy)):
            value = _number(record["values"][i])
            out.append(dict(zip(FIELDS, (fy, month, mineral, record["province"],
                "" if value is None else value, "" if printed is None else printed,
                "year_total" if printed is not None else "absent-in-edition", canonical, sector,
                "T2" if timing and canonical else "T4" if canonical else "",
                "physical_production" if canonical else "", basis, os.path.basename(path), unit,
                _status(record["values"][i]), sheet, record["source_row"], header,
                _publication_status(header), group_id, "" if national_values[i] is None else national_values[i],
                expected, "pass" if closed else "fail", "pass" if annual_ok else "fail",
                "complete" if closed else "incomplete", "yes" if timing else "no", quality))))
    control = dict(zip(CONTROL_FIELDS, (fy, mineral, canonical, os.path.basename(path), sheet,
        header, _publication_status(header), group_id, expected, sum(national_values) if complete else "",
        sum(_number(r["values"][i]) for r in admitted for i in range(12)) if complete else "",
        "" if residual is None else residual, "yes" if closed else "no", "complete" if closed else "incomplete",
        "yes" if timing else "no", quality, ";".join("" if v is None else str(v) for v in national_values),
        "no", "witness only")))
    return out, bad, control

def parse(path):
    rows, bad, groups = [], [], []
    for group in (_xlsx_groups(path) if str(path).endswith(".xlsx") else _pdf_groups(path)):
        parsed, failed, control = _assemble(path, group)
        rows.extend(parsed); bad.extend(failed); groups.append(control)
    return rows, bad, groups

def validate_export(rows):
    """Validate actual final records against printed national monthly cells."""
    keys = [(r["fiscal_year_start"], r["mineral"], r["province"], r["month"]) for r in rows]
    if len(keys) != len(set(keys)):
        raise ValueError("duplicate final mineral/reporting-area/month key")
    grouped = collections.defaultdict(list)
    for row in rows:
        grouped[(row["fiscal_year_start"], row["mineral"], row["month"])].append(row)
    for key, members in grouped.items():
        if members[0]["national_coverage"] != "complete":
            if any(r["timing_usable"] == "yes" for r in members):
                raise ValueError("incomplete source group admitted for timing: %s" % (key,))
            continue
        expected = set(members[0]["source_expected_provinces"].split(";"))
        if {r["province"] for r in members} != expected:
            raise ValueError("final reporting areas differ from source group: %s" % (key,))
        controls = {float(r["source_national_month_quantity"]) for r in members}
        if len(controls) != 1 or abs(sum(float(r["quantity"]) for r in members) - controls.pop()) > GROUP_TOLERANCE:
            raise ValueError("final exported provinces fail national monthly control: %s" % (key,))

def collect():
    files = sorted(glob.glob(asset_path(RAW, "pbs_minerals_*.pdf")) + glob.glob(asset_path(RAW, "pbs_minerals_*.xlsx")))
    all_rows, bad, controls = [], [], []
    for path in files:
        rows, failed, groups = parse(path)
        all_rows.extend(rows); bad.extend(failed); controls.extend(groups)
    candidates = collections.defaultdict(list)
    for control in controls:
        if control["canonical_mineral"]:
            candidates[(control["fiscal_year_start"], control["canonical_mineral"])].append(control)
        else:
            control["selection_reason"] = "excluded: outside the four supported inventory-driver minerals; source retained"
    chosen = set()
    for key, groups in sorted(candidates.items()):
        winner = min(groups, key=lambda g: (g["source_file"] != COMPILED, g["closed"] != "yes",
                                          g["source_file"], g["group_id"]))
        winner["selected"] = "yes"
        winner["selection_reason"] = ("printed-year block in held compiled workbook" if winner["source_file"] == COMPILED
                                      else "held annual PDF/XLSX edition")
        chosen.add(winner["group_id"])
    rows = sorted((r for r in all_rows if r["group_id"] in chosen),
                  key=lambda r: (r["fiscal_year_start"], r["mineral"], r["province"], r["month"]))
    validate_export(rows)
    return files, rows, bad, controls

def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__); parser.add_argument("--check", action="store_true")
    args = parser.parse_args(argv)
    files, rows, bad, controls = collect()
    years = {r["fiscal_year_start"] for r in rows}; selected = [g for g in controls if g["selected"] == "yes"]
    summary = [{"files_parsed": len(files), "fiscal_years_held": len(years),
        "fiscal_years_absent": ";".join("FY%d-%02d" % (fy, (fy + 1) % 100) for fy in WINDOW if fy not in years) or "none",
        "month_rows": len(rows), "rows_quarantined": len(bad), "mineral_years": len(selected),
        "mineral_years_closing_group_control": sum(g["closed"] == "yes" for g in selected),
        "canonical_minerals_found": ";".join(sorted({r["canonical_mineral"] for r in rows if r["canonical_mineral"]})),
        "canonical_year_coverage": ";".join("%s:%d/%d" % (m, len({r["fiscal_year_start"] for r in rows if r["canonical_mineral"] == m}), len(years))
                                             for m in sorted({r["canonical_mineral"] for r in rows if r["canonical_mineral"]})),
        "wired_into_any_sector": "fugitives upstream validated timing; limestone cross-check"}]
    results = [write_or_check(OUT, FIELDS, rows, args.check, banner=[
        "PBS natural gas, crude oil, coal and limestone production at the printed fiscal-year grain; native units are per row.",
        "One explicitly selected source group per mineral-year; historical blocks and duplicate worksheets are not added.",
        "Unknown cells remain empty. Complete final provincial records close on printed national months.",
        "FY2015-16 is held inside the compiled workbook. Suspect FY2021-22 gas/oil observations are excluded from timing."]),
        write_or_check(QUARANTINE, BAD_FIELDS, bad, args.check), write_or_check(CONTROLS, CONTROL_FIELDS, controls, args.check),
        write_or_check(SUMMARY, list(summary[0]), summary, args.check)]
    if not args.check:
        validate_export(read_rows(OUT))
    else:
        for path, result in zip((OUT, QUARANTINE, CONTROLS, SUMMARY), results):
            if not result:
                print("checked: " + os.path.relpath(path, ROOT))
    print("PBS minerals: %d records, %d fiscal years; final exported controls validated" % (len(rows), len(years)))
    return int(any(results)) if args.check else 0

if __name__ == "__main__":
    raise SystemExit(main())
