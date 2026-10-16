"""Reproduce the FY2024-25 inputs and separately scoped FY2025-26 witnesses.

Offline: all primary bytes must already be retained by fetch_release_extension.py.
The final exported tables, not temporary accumulators, are independently checked
by tests/test_release_extension_inputs.py. --check compares without writing.
"""
from __future__ import annotations

import argparse
import collections
import hashlib
import importlib.util
import os
from pathlib import Path
import re
import subprocess
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from pkinv.assets import asset_path
from pkinv.core.io import read_rows, write_or_check
from pkinv.sectors.transport import ocac

RAW = ROOT / "data/raw"
AGRICULTURE = RAW / "finance_division/survey_2025_26/2_Agriculture.pdf"
QIM = RAW / "pbs/qim/extension_2026"
OCAC = RAW / "ocac/extension_2026"
NEPRA = ROOT / "sources/Power Plants/NEPRA State of Industry Report 2025.pdf"
BASELINE = ROOT / "data/baselines/rc67_integrated_campaign"


def qim_banner():
    return ["PBS physical monthly output, ten fiscal years July2015-June2025.",
            "Printed monthly/annual quantity controls retained; index rows are separately typed.",
            "Product output is a process timing proxy (T2), not observed process emissions."]


def text(path):
    return subprocess.check_output(["pdftotext", "-layout", str(path), "-"]).decode("utf-8", "replace")


def relative(path):
    return Path(path).relative_to(ROOT).as_posix()


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def csv_at(path):
    return read_rows(asset_path(ROOT, path))


def months(start):
    return ["%04d%02d" % (start + (m < 7), m) for m in list(range(7, 13)) + list(range(1, 7))]


def fiscal(start):
    return "FY%d-%02d" % (start, (start + 1) % 100)


def agriculture_tables():
    pages = text(AGRICULTURE).split("\f")
    tables = {}
    for number, title in (("2.3", "AREA UNDER IMPORTANT CROPS"),
                          ("2.4", "PRODUCTION OF IMPORTANT CROPS"),
                          ("2.9", "FERTILIZER OFFTAKE AND IMPORTS")):
        selected = [(i, page) for i, page in enumerate(pages, 1) if title in page]
        if len(selected) != 1:
            raise ValueError("Agriculture table %s is ambiguous" % number)
        page_no, page = selected[0]
        values = {}
        for line in page.splitlines():
            match = re.match(r"^\s*(20\d\d)-(\d\d)\s+(?:P\s+)?(.*)$", line)
            if not match:
                continue
            year = int(match.group(1))
            cells = match.group(3).split()
            if not all(re.fullmatch(r"[\d,.]+", c) for c in cells):
                continue
            values[year + 1] = [float(c.replace(",", "")) for c in cells]
        tables[number] = (page_no, values)
    return tables


def agriculture_rows():
    tables = agriculture_tables()
    area_columns = {"wheat": 0, "rice": 1, "maize": 4, "sugarcane": 8, "cotton": 11}
    production_columns = {"wheat": 0, "rice": 1, "maize": 4, "sugarcane": 8, "cotton": 12}
    source = "Pakistan Economic Survey 2025-26 %s" % relative(AGRICULTURE)
    rows, provisional = [], []
    for end in range(2016, 2027):
        current = []
        for kind, columns, table in (("area", area_columns, "2.3"),
                                     ("production", production_columns, "2.4")):
            page, values = tables[table]
            cells = values[end]
            if len(cells) != (13 if kind == "area" else 14):
                raise ValueError("Agriculture table %s/%d column count differs" % (table, end))
            for crop, index in columns.items():
                cotton = crop == "cotton" and kind == "production"
                current.append({"fiscal_end_year": end, "fiscal_label": fiscal(end - 1),
                    "item": "cotton_bales_000" if cotton else crop + ("_area_ha" if kind == "area" else "_production_t"),
                    "value": int(cells[index] * (1 if cotton else 1000)),
                    "unit": "000 bales" if cotton else "ha" if kind == "area" else "t",
                    "source_table": source + " appendix Table%s PDFp%d" % (table, page),
                    "basis": "provisional_fiscal_crop" if end == 2026 else "observed_fiscal_revised"})
        if end <= 2025:
            page, values = tables["2.9"]
            current.append({"fiscal_end_year": end, "fiscal_label": fiscal(end - 1),
                "item": "nitrogen_N_offtake_t", "value": int(values[end][0] * 1000), "unit": "t N",
                "source_table": source + " appendix Table2.9 PDFp%d" % page,
                "basis": "observed_fiscal_full_year"})
            rows.extend(current)
        else:
            provisional.extend(current)
    return rows, provisional


def qim_records(path, start):
    import openpyxl
    book = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheet = book.active
    header = next(sheet.iter_rows(min_row=5, max_row=5, values_only=True))
    if [v.strftime("%Y%m") for v in header[4:16]] != months(start):
        raise ValueError("QIM filename and printed month columns disagree")
    out = []
    for number, cells in enumerate(sheet.iter_rows(values_only=True), 1):
        serial = str(cells[0] or "").strip()
        if not re.fullmatch(r"\d+(?:-\d+)?", serial) or not cells[1]:
            continue
        # Group production rows 1-3 are a physical sum. Other ranges may be
        # indices despite carrying a unit: the printed annual SUM gate distinguishes them.
        values, annual = cells[4:16], cells[16]
        if not all(isinstance(v, (int, float)) for v in values + (annual,)):
            continue
        if abs(sum(values) - annual) > 2:
            continue
        item = str(cells[1]).strip()
        for month, value in zip(months(start), values):
            out.append({"fiscal_year_start": start, "month": month, "item_no": serial,
                "item": item, "unit": str(cells[2] or "").strip(), "weight": cells[3] or "",
                "row_kind": "quantity", "annual_control": "sum", "quantity": float(value),
                "annual_printed": float(annual),
                "canonical_item": {"83": "cement", "85": "billets_ingots", "82": "glass_plates_sheets", "1-3": "sugar"}.get(serial, ""),
                "consuming_sector": {"83": "ippu", "85": "industry", "82": "industry", "1-3": "power"}.get(serial, ""),
                "temporal_evidence_class": "TE2" if serial in ("83", "85", "82", "1-3") else "",
                "timing_target": "physical_output", "basis": "Own-year production; process timing proxy, not observed emissions",
                "source_file": relative(path)})
    book.close()
    required = {"83", "85", "58", "82", "1-3"}
    if not required <= {r["item_no"] for r in out}:
        raise ValueError("QIM omitted a required physical row")
    return out


def power_products():
    """Append verified plant quantities; derive net energy explicitly by fuel.

    CPPA/GENCO remains the common boundary. KE quantities stay in the plant ledger
    and a separate coverage table; national generation controls include KE separately.
    """
    source_rows = csv_at("data/sectors/power/activity/plant_fuel_2024_25.csv")
    ledger = csv_at("data/activity/power/plant_fuel_ledger.csv")
    fields = list(ledger[0])
    ledger = [r for r in ledger if int(r["fiscal_year"]) != 2024]
    ef = {(r["category"], r["item"]): float(r["value"]) for r in csv_at("data/emission_factors.csv") if r["value"]}
    # Identical physical conversion conventions to existing residential/industry gas:
    # 1000Btu/scf (gross) and NCV/GCV0.90. MMBTU is gross billed energy.
    factor = {"Gas": (0.00105505585262 * .90), "RLNG": (0.00105505585262 * .90),
              "RFO": ef[("industry_ncv", "Residual Fuel Oil")],
              "HSD": ef[("industry_ncv", "Gas/Diesel Oil")]}
    # NCV defaults are TJ/Gg; kt has the same numerical magnitude as Gg.
    energies = collections.defaultdict(float)
    generation = collections.defaultdict(float)
    new_ledger, records = [], []
    station_intensities = collections.defaultdict(list)
    for r in source_rows:
        q, unit, fuel = float(r["quantity"]), r["quantity_unit"], r["fuel"]
        energy = ""
        if fuel in ("Gas", "RLNG"):
            energy = q * factor[fuel] * (1000 if unit == "MMCFT" else 1)
        elif fuel == "RFO":
            energy = q * factor[fuel]
        elif fuel == "HSD":
            # 000litres ×0.84kg/L /1000 = Gg. Density is an explicit transfer assumption.
            energy = q * .84 / 1000 * factor[fuel]
        item = {k: r.get(k, "") for k in fields}
        item["energy_TJ"] = energy
        item["implied_heat_rate_kJ_per_kWh"] = (float(energy) * 1000 / float(r["net_generation_GWh"])
                                                   if energy != "" and float(r["net_generation_GWh"]) else "")
        item["source_table"] += "; " + r["estimate_status"] + ("; " + r["note"] if r["note"] else "")
        if energy != "":
            item["source_table"] += "; energy_TJ is derived NCV energy: " + (
                "gross MMBTU or 1000 Btu/scf volume times NCV/GCV 0.90"
                if fuel in ("Gas", "RLNG") else
                "mass times IPCC NCV" if fuel == "RFO" else
                "000 litres times 0.84 kg/L and IPCC diesel NCV")
        new_ledger.append(item)
        system = "K-Electric" if r["system"] == "K-Electric" else "CPPA-G"
        generation[(system, fuel)] += float(r["net_generation_GWh"])
        if energy != "":
            energies[(system, fuel)] += float(energy)
            if float(r["net_generation_GWh"]) > 0:
                station_intensities[(system, fuel)].append(float(energy) / float(r["net_generation_GWh"]))
    # Read the final FY column from the printed Table17, independently of the station sums.
    pages = text(NEPRA).split("\f")
    page = next(p for p in pages if "Thermal Electricity Generation by Sector and by Fuel" in p and "GENCOs (I" in p)
    control_values = []
    for line in page.splitlines():
        if re.search(r"Generation on (?:Gas|RLNG|RFO|HSD|Coal) \(GWh", line):
            vals = re.findall(r"\d[\d,]*(?:\s?\.\s?\d+)?", line.split("(GWh", 1)[1])
            control_values.append(float(vals[-1].replace(",", "").replace(" ", "")))
    if len(control_values) != 15:
        raise ValueError("Table17 must supply five fuels × CPPA,KE,total")
    controls = {(system, fuel): control_values[3 * i + j]
                for i, fuel in enumerate(("Gas", "RLNG", "RFO", "HSD", "Coal"))
                for j, system in enumerate(("CPPA-G", "K-Electric"))}
    # An explicitly recorded fleet-intensity estimate fills control gaps; no named
    # plant or measured fuel quantity is manufactured. Coal booking is separate.
    for (system, fuel), amount in controls.items():
        placed = generation.get((system, fuel), 0)
        residual = amount - placed
        if residual < -.1:
            raise ValueError("Plant generation exceeds %s/%s control" % (system, fuel))
        inferred = 0.0
        gap_low = gap_high = 0.0
        if residual > .1 and fuel != "Coal":
            if placed <= 0:
                raise ValueError("No supported intensity for missing fuel/control")
            inferred = residual / placed * energies[(system, fuel)]
            intensity = station_intensities[(system, fuel)]
            gap_low, gap_high = residual * min(intensity), residual * max(intensity)
            energies[(system, fuel)] += inferred
        records.append({"fiscal_year_start": 2024, "system": system, "fuel": fuel,
            "control_generation_GWh": amount, "station_generation_GWh": round(placed, 6),
            "unresolved_generation_GWh": round(max(0, residual), 6), "estimated_gap_energy_TJ": inferred,
            "gap_energy_low_TJ": gap_low, "gap_energy_high_TJ": gap_high,
            "station_plus_gap_energy_TJ": "" if fuel == "Coal" else energies[(system, fuel)],
            "estimate_status": ("not_estimated_missing_fuel" if fuel == "Coal" and residual > .1 else
                                "derived_fleet_intensity_gap" if inferred else "physical_station_fuel"),
            "energy_basis": "NCV; gas assumed 1000Btu/scf gross ×0.90; diesel density0.84kg/L; IPCC oilNCV",
            "source": relative(NEPRA) + " Tables17/18/21/26",
            "note": "Coal energy uses separate rank/NCV booking. KE is coverage alternative. Positive zero-generation fuel is retained."})
    new_generation = [{"fiscal_year": 2024, "fuel": f, "system": s, "gwh_net": q,
                       "source": "NEPRA State of Industry Report 2025 Table17 (net); " + relative(NEPRA)}
                      for (s, f), q in sorted(controls.items())]
    old_fuel = csv_at("data/activity/power/fuel_consumed.csv")
    fuel_rows = [r for r in old_fuel if int(r["fiscal_year"]) != 2024]
    # Coal TOE here is a source-control placeholder derived from physical station
    # ranks, not the reference booking; the accounting producer owns that booking.
    coal = sum(float(r["quantity"]) * (11.9 if r["fuel_origin"] == "Thar lignite" or r["plant"] == "Lucky Electric" else 25.8)
               for r in source_rows if r["fuel"] == "Coal" and r["system"] != "K-Electric")
    net = {"Coal": coal, "Gas": energies[("CPPA-G", "Gas")] + energies[("CPPA-G", "RLNG")],
           "Furnace Oil": energies[("CPPA-G", "RFO")], "Diesel Oil": energies[("CPPA-G", "HSD")]}
    for fuel, tj in net.items():
        fuel_rows.append({"fiscal_year": 2024, "fuel": fuel, "toe": tj / .041868,
            "source": "NEPRA SoIR2025 station consumption Tables18/21, CPPA/GENCO boundary; "
                      "physical fuel converted on NCV; unresolved Table17 generation at own-year fleet intensity; "
                      "see pakemis_power_extension_controls.csv. Coal reference is separately generated rank/NCV booking."})
    import ingest_nepra_bagasse as bagasse_ingest
    bagasse_text = re.sub(r"(\d)\s+\.(\d)", r"\1.\2", "\n".join(pages))
    bagasse_rows, totals = bagasse_ingest.parse(bagasse_text.splitlines(), years=list(range(2020, 2025)))
    own_bagasse = [r for r in bagasse_rows if r["fiscal_year"] == 2024]
    if totals is None or abs(sum(r["generation_GWh"] for r in own_bagasse) - totals[-1]) > .02:
        raise ValueError("FY2025 bagasse stations do not close on the printed total")
    for row in own_bagasse:
        row["source_table"] = "NEPRA SoIR2025 Table8 E3; " + relative(NEPRA)
    return ledger + new_ledger, fuel_rows, new_generation, records, own_bagasse


def ocac_extension():
    annual_path = OCAC / "Jul-Jun-2025-26-ENE-Pak-2.pdf"
    annual = ocac.parse_annual(str(annual_path))
    rows, revisions, checks = [], [], []
    old = {(r["section"], r["sector_row"], r["product"]): float(r["tonnes"])
           for r in csv_at("data/activity/transport/ocac_sectoral_annual.csv")
           if int(r["fiscal_year_start"]) == 2024}
    for section, name, quantities in annual["rows"]:
        if (section, name, "TOTAL") in old and "TOTAL_PRIOR" in quantities:
            before, after = old[(section, name, "TOTAL")], quantities["TOTAL_PRIOR"]
            revisions.append({"period": "FY2024-25", "section": section, "customer_class": name,
                "old_tonnes": before, "revised_tonnes": after, "delta_tonnes": after - before,
                "proportional_composition_multiplier": after / before if before else "",
                "old_source": "sources/Aviation/ocac_annual/OCAC_FY2024_25.pdf",
                "new_source": relative(annual_path),
                "status": "published_aggregate_revision_product_split_unresolved",
                "classification": "", "material": "yes" if before and abs(after - before) / before >= 0.005 else "no",
                "first_print_grand_total_gap_tonnes": "",
                "note": "Core retains complete2025product edition. Multiplier is an explicit unchanged-composition sensitivity, not observed revised product cells."})
        for product, value in quantities.items():
            if product in ("TOTAL", "TOTAL_PRIOR"):
                continue
            rows.append({"period": "FY2025-26", "month": "", "area": "PAKISTAN", "sector": section,
                "activity": name + ":" + product, "quantity": value, "unit": "t",
                "status": "provisional_full_year_sales", "source": relative(annual_path),
                "boundary": "Customer-class fuel sales; class totals overlap member rows; exports/bunkers outside domestic totals",
                "temporal_support": "annual sales", "spatial_support": "national"})
    # Classify the FY2024-25 overlap with the same rule as the historical editions.
    total_delta = next((r["delta_tonnes"] for r in revisions if r["customer_class"] == "TOTAL_INTERNAL"), 0)
    for r in revisions:
        r["classification"] = ("unchanged_printed_rounding" if abs(r["delta_tonnes"]) <= 2 else
                               "reclassification_between_customer_classes" if abs(total_delta) <= 2 and
                               r["customer_class"] != "TOTAL_INTERNAL" else "level_revision")
    revisions = ocac_overlap_history() + revisions
    by_month = collections.defaultdict(float)
    for path in sorted(OCAC.glob("ocac_sector_province_*.pdf")):
        year, month = re.search(r"_(\d{4})_(\d{2})\.pdf$", path.name).groups()
        period = year + month
        for area, parsed in ocac.parse_monthly(str(path)):
            for warning in parsed["warnings"]:
                checks.append({"period": "FY2025-26", "month": period, "area": area,
                    "check": "publisher_table_row", "residual": "", "unit": "t",
                    "status": "retained_warning", "note": warning, "source": relative(path)})
            for section, name, quantities in parsed["rows"]:
                for product, value in quantities.items():
                    if product in ("TOTAL", "TOTAL_PRIOR"):
                        continue
                    rows.append({"period": "FY2025-26", "month": period, "area": area, "sector": section,
                        "activity": name + ":" + product, "quantity": value, "unit": "t",
                        "status": "provisional_monthly_sales", "source": relative(path),
                        "boundary": "Karachi is contained in Sindh; customer classes and totals overlap. Sales are not combustion observations.",
                        "temporal_support": "own-year monthly sales", "spatial_support": "own-year provincial sales; Karachi nonadditive"})
                    if area == "PAKISTAN":
                        by_month[(section, name, product)] += value
    for section, name, quantities in annual["rows"]:
        if name not in ocac.COMMIT_ROWS:
            continue
        for product, value in quantities.items():
            key = (section, name, product)
            if key not in by_month:
                continue
            residual = by_month[key] - value
            checks.append({"period": "FY2025-26", "month": "", "area": "PAKISTAN",
                "check": name + ":" + product + " monthly-minus-annual", "residual": residual,
                "unit": "t", "status": "closed_printed_precision" if abs(residual) <= 12 else "source_edition_revision",
                "note": "Unadjusted monthly sales; emissions use annual quantities and explicitly normalised monthly shape.",
                "source": relative(annual_path)})
    return rows, revisions, checks


NFDC_JULY_2026 = RAW / "nfdc/monthly_reviews_fy2026/nfdc_mfr_2026_07.pdf"


def nfdc_fy2026_nitrogen():
    """FY2025-26 monthly nitrogen offtake from the NFDC July 2026 review, Annex-3.

    Annex-3 ("month-wise comparative nutrient offtake during fiscal year 2026-27 compared
    with previous years") prints July-June columns for FY2022-23..FY2026-27; the FY2026-27
    column is flagged "* July onward are average of previous year" and is never read. The
    FY2025-26 column is the first full twelve-month nitrogen record for that year. Two gates
    against independent prints: the FY2024-25 column must equal the held annual review's
    3,387,074 t N control (Survey 2025-26 Table 2.9 prints 3,387 kt), and FY2025-26
    July-March must reproduce Survey Table 2.9's provisional 3,035 kt. Offtake is product
    sales, not field application.
    """
    if not NFDC_JULY_2026.exists():
        return [], None
    # pdftotext warns about this file's optional-content table on stderr; the text is complete.
    lines = subprocess.check_output(["pdftotext", "-layout", str(NFDC_JULY_2026), "-"],
                                    stderr=subprocess.DEVNULL).decode("utf-8", "replace").splitlines()
    start = next(i for i, line in enumerate(lines) if "FISCAL YEAR 2026-27 COMPARED" in line)
    header = next(line for line in lines[start:start + 6] if line.strip().startswith("YEAR"))
    labels = re.findall(r"FY\s*(\d{4})-(\d{2})", header)
    column = [int(a) for a, _ in labels].index(2025)
    control = [int(a) for a, _ in labels].index(2024)
    block = next(i for i in range(start, len(lines)) if lines[i].strip() == "NITROGEN")
    names = ("July", "August", "September", "October", "November", "December",
             "January", "February", "March", "April", "May", "June")
    values, prior = {}, {}
    for line, name, month in zip(lines[block + 1:block + 13], names, months(2025)):
        cells = line.split()
        if cells[0] != name:
            raise ValueError("NFDC Annex-3 month order differs at %r" % line)
        numbers = [float(c.replace(",", "")) for c in cells[1:len(labels) + 1]]
        values[month], prior[month] = numbers[column], numbers[control]
    total_cells = lines[block + 13].split()
    if total_cells[0] != "Total":
        raise ValueError("NFDC Annex-3 nitrogen block has no Total row")
    printed_total = float(total_cells[1 + column].replace(",", ""))
    if abs(sum(values.values()) - printed_total) > 2:
        raise ValueError("NFDC Annex-3 FY2025-26 nitrogen months fail the printed Total")
    if abs(sum(prior.values()) - 3387074) > 2:
        raise ValueError("NFDC Annex-3 FY2024-25 column differs from the held AFR2024-25 control")
    july_march = sum(values[m] for m in months(2025)[:9])
    if abs(july_march / 1000 - 3035) > 1:
        raise ValueError("NFDC Annex-3 July-March differs from Survey 2025-26 Table 2.9")
    source = relative(NFDC_JULY_2026) + " Annex-3 nitrogen, FY2025-26 column"
    rows = [{"period": fiscal(2025), "month": month, "area": "PAKISTAN", "sector": "agriculture",
             "activity": "nitrogen_N_offtake_t", "quantity": quantity, "unit": "t N",
             "status": "provisional_monthly_offtake", "source": source,
             "boundary": "Fertilizer product offtake (dispatch to market), not field application or emission timing",
             "temporal_support": "own-year monthly offtake", "spatial_support": "national only"}
            for month, quantity in sorted(values.items())]
    return rows, printed_total


def ocac_overlap_history():
    """Every overlapping OCAC year between consecutive annual editions, classified.

    Edition Y prints its own year (TOTAL) and the previous year (TOTAL_PRIOR) per customer
    class. A difference is: unchanged (|delta| <= 2 t, printed rounding); a correction of an
    internally non-additive first print (the earlier edition's classes fail its own printed
    grand total by the same amount the later edition restores); a reclassification between
    classes (the grand total unchanged); or a level revision (the grand total moves).
    The core keeps each year's own first edition; nothing here replaces a product cell.
    """
    core = csv_at("data/activity/transport/ocac_sectoral_annual.csv")
    first, restated, files = {}, {}, {}
    for r in core:
        year = int(r["fiscal_year_start"])
        if r["product"] == "TOTAL":
            first[(year, r["section"], r["sector_row"])] = float(r["tonnes"])
            files[year] = r["source_file"]
        elif r["product"] == "TOTAL_PRIOR":
            restated[(year - 1, r["section"], r["sector_row"])] = float(r["tonnes"])
    classes = ("DOMESTIC", "AGRICULTURE", "Total Transport", "Total Industry", "Total Power", "Total Government")
    # rc68: the core ingest now books a publisher's correction of a non-additive first print
    # (scripts/ingest_ocac_sectoral.py DECLARED_PUBLISHER_CORRECTIONS). The comparison here is
    # made against the first print as printed, so the correction stays visible as applied.
    import importlib.util as _ilu
    _spec = _ilu.spec_from_file_location(
        "ingest_ocac_sectoral", os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                             "ingest_ocac_sectoral.py"))
    _ocac_ingest = _ilu.module_from_spec(_spec)
    _spec.loader.exec_module(_ocac_ingest)
    applied_corrections = {}
    for (cy, csec, crow, cprod), rec in _ocac_ingest.DECLARED_PUBLISHER_CORRECTIONS.items():
        aggregate = {"TRANSPORT": "Total Transport"}.get(csec)
        for row in (crow,) + ((aggregate,) if aggregate else ()):
            applied_corrections[(cy, csec, row)] = rec["gap_tonnes"]
    for key, gap in list(applied_corrections.items()):
        if key in first:
            first[key] = first[key] - gap
    out = []
    for (year, section, name), after in sorted(restated.items()):
        before = first.get((year, section, name))
        if before is None:
            continue
        total_key = (year, "TOTALS", "TOTAL_INTERNAL")
        first_sum = sum(v for (y, s, n), v in first.items() if y == year and n in classes)
        first_gap = first.get(total_key, 0) - first_sum
        total_delta = restated.get(total_key, 0) - first.get(total_key, 0)
        delta = after - before
        booked = (year, section, name) in applied_corrections
        if abs(delta) <= 2:
            kind = "unchanged_printed_rounding"
        elif abs(first_gap) > 2 and abs(abs(delta) - abs(first_gap)) <= 2 and abs(total_delta) <= 2:
            kind = "correction_of_non_additive_first_print"
        elif abs(total_delta) <= 2 and name != "TOTAL_INTERNAL":
            kind = "reclassification_between_customer_classes"
        else:
            kind = "level_revision"
        out.append({"period": fiscal(year), "section": section, "customer_class": name,
            "old_tonnes": before, "revised_tonnes": after, "delta_tonnes": delta,
            "proportional_composition_multiplier": after / before if before else "",
            "old_source": "sources/Aviation/ocac_annual/" + files[year],
            "new_source": "sources/Aviation/ocac_annual/" + files.get(year + 1, ""),
            "status": ("applied_" + kind) if booked else ("retained_first_print_" + kind),
            "classification": kind,
            "material": "yes" if before and abs(delta) / before >= 0.005 else "no",
            "first_print_grand_total_gap_tonnes": first_gap,
            "note": ("BOOKED since rc68: the first print's classes fall short of its own printed "
                     "product total by exactly this amount and the next edition restores it to this "
                     "row, so the core carries the corrected value (scripts/ingest_ocac_sectoral.py)."
                     if booked else
                     "Core keeps each year's own first edition; the later edition prints only the class "
                     "total, so no product cell is replaced.")})
    return out


def provisional_emissions(activity):
    """Supported source-level estimates only. No national aggregate is produced."""
    ef = {(r["category"], r["item"]): float(r["value"]) for r in csv_at("data/emission_factors.csv") if r["value"]}
    mapping = {"MS": ("Motor Gasoline", 44.3), "HOBC": ("Motor Gasoline", 44.3),
               "HSD": ("Gas/Diesel Oil", 43), "LDO": ("Gas/Diesel Oil", 43),
               "FO": ("Residual Fuel Oil", 40.4), "KERO": ("Other Kerosene", 43.8)}
    annual = [r for r in activity if r["month"] == "" and r.get("area") == "PAKISTAN"]
    out = []
    for r in annual:
        if ":" not in r["activity"]:
            continue
        name, product = r["activity"].rsplit(":", 1)
        factors = {}
        method = ""
        if name == "Road" and product in ("MS", "HOBC", "HSD"):
            fuel, ncv = mapping[product]
            factors = {g: ef[("transport_fuel_" + g.lower(), fuel)] * ncv / 1000
                       for g in ("CO2", "CH4", "N2O")}
            method = "Retail Road customer class; road-factor application before unresolved off-road partition; CNG absent"
        elif name == "Total Industry" and product in mapping:
            fuel, ncv = mapping[product]
            factors = {"CO2": ef[("industry_combustion_co2", fuel)] * ncv / 1000,
                       "CH4": ef[("industry_combustion_ch4", "Liquid Fuels")] * ncv / 1000,
                       "N2O": ef[("industry_combustion_n2o", "Liquid Fuels")] * ncv / 1000}
            method = "Industrial liquid combustion only; no gas or coal inferred"
        elif name in ("Domestic Avi-PIA", "Domestic Avi-Others") and product == "JP-1":
            factors = {"CO2": 3160, "SO2": ef[("aviation_fuel_ef", "SO2")],
                       "PM2.5": ef[("aviation_fuel_ef", "PM2.5")]}
            factors["BC"] = factors["PM2.5"] * ef[("aviation_fuel_ef", "BC_fraction")]
            method = "Resident-airline customer-class uplift; domestic/international flight split unresolved; carbon and transferred fuel-mass pollutant endpoint; no LTO NOx estimate"
        if not factors:
            continue
        for species, kg_per_t in factors.items():
            out.append({"period": "FY2025-26", "month": "", "area": "PAKISTAN",
                "source_process": r["activity"], "species": species, "emissions_Gg": float(r["quantity"]) * kg_per_t / 1e6,
                "mass_convention": "native pollutant mass", "status": "provisional_source_estimate",
                "configuration": "supported_source_extension", "activity_tonnes": r["quantity"],
                "factor_kg_per_tonne": kg_per_t, "method": method, "source": r["source"],
                "coverage": "Partial source panel; rows must not be called a common-scope national total",
                "temporal_support": "Annual; monthly allocation supplied separately from own-year sales",
                "spatial_support": "National; provincial sales are allocation proxies, not release locations"})
    # Billets/ingots uses the inventory's admitted electric-route process boundary.
    billets = [r for r in activity if r["month"] and r["activity"] == "Billets/Ingots" and r.get("area") == "PAKISTAN"]
    if billets:
        quantity = sum(float(r["quantity"]) for r in billets)
        out.append({"period": "FY2025-26", "month": "", "area": "PAKISTAN",
            "source_process": "Billets/Ingots electric-route process", "species": "CO2",
            "emissions_Gg": quantity * .08 / 1000, "mass_convention": "native CO2 mass",
            "status": "provisional_source_estimate", "configuration": "supported_source_extension",
            "activity_tonnes": quantity, "factor_kg_per_tonne": 80,
            "method": "Existing IPCC electric-route80kgCO2/t billets; route transfer is a conditional assumption, rolling outputs excluded",
            "source": billets[0]["source"], "coverage": "Electric-route process only; no power generation or rolling added",
            "temporal_support": "Own-year monthly billet output", "spatial_support": "National"})
    return out


def availability_rows():
    sources = {
        "power": (relative(NEPRA), "NEPRA SoIR2025 Tables17/18/21/26", "physical station fuel/net generation; CPPA/GENCO reference", "FY2026 annual physical station fuel not established; monthly generation filings alone insufficient"),
        "transport": ("sources/Aviation/ocac_annual/OCAC_FY2024_25.pdf", "OCAC full FY2025 edition; FY2026 overlap audit", "fuel/product/customer sales; CNG from held OGRA2024-25", "FullFY2026 roadliquid sales supported; annual CNG missing"),
        "industry": ("data/raw/finance_division/survey_2025_26/14_Energy.pdf", "Survey2025-26/OGRA2024-25/OCAC", "FullFY2025 coal controls, industrial gas and liquids", "FY2026 liquids supported; coal/gas observations July-March only"),
        "agriculture": (relative(AGRICULTURE), "Survey2025-26 appendix2.3/2.4/2.9; NFDC AFR2024-25", "FullFY2025 crops/nutrients; census-anchored livestock reconstruction", "FY2026 crops provisional; full-year nitrogen offtake held (product offtake, not application); livestock projection only"),
        "residential": ("sources/HIES/HIES-2024-25-Report-Final-2.pdf", "HIES2024-25 quantities/OGRA2024-25 full-year gas", "S survey quantities; L structural alternative; survey reference is not twelve observed months", "No new annual household survey; gas/LPG fullFY2026 not established"),
        "waste": ("data/sectors/waste/activity/population_district.csv", "Population projection, held HIES2024-25 treatment evidence", "Advance disposal-history state; annual management shares not observed", "Population/state continuation possible; no new management observation claimed"),
        "ippu": (relative(QIM / "Table-3-2024-25-6.xlsx"), "PBS June2026 physical production tables; NFDC AFR2024-25", "FullFY2025 cement/billets; fertilizer physical products for existing ammonia reconstruction", "FY2026 billets/soda/cement output supported; no physical clinker/ammonia row"),
        "fugitives": ("data/raw/ogra/state_of_industry/ogra_petroleum_industry_report_2024_25.pdf", "OGRA2024-25; Survey2025-26 ManufacturingTable3.11", "FullFY2025 production/deliveries/infrastructure; monthly PBS absent givesT4", "FY2026 annual gas/production/infrastructure not established; Jul-Mar kept partial"),
        "railways": ("sources/Railways/yearbook2024-25.pdf", "Railway YearBook2024-25 operator fuel, provisional pending audit", "FullFY2025 traction and other fuel quantities", "FY2026 OCAC sales cannot establish operator traction/non-traction fuel boundary"),
        "aviation": ("data/raw/pcaa/extension_2025/major-traffic-flows-by-airports-during-the-year-2024-25.pdf", "PCAA airport trafficFY2024-25; OCAC annual/monthlyFY2025 andFY2026", "Observed own-year airport domestic/international/noncommercial movements; separate fuel customer classes", "FY2026 resident-airline customer-class uplift supports separate partial-source carbon and transferred fuel-mass endpoints; domestic/international flights and own-year LTO allocation unresolved"),
    }
    candidate_months = {
        "power": "No FY2025-26 full-year physical station-fuel panel established; latest is2024-07 through2025-06",
        "transport": "Liquid customer sales2025-07 through2026-06 (12months); gas2025-07 through2026-03 (9months)",
        "industry": "Liquid customer sales2025-07 through2026-06 (12months); coal/gas2025-07 through2026-03 (9months)",
        "agriculture": "FY2025-26 provisional crop-season totals; nitrogen offtake2025-07 through2026-06 (12months, NFDC July2026 review Annex-3, provisional); livestock has no FY2025-26 observation",
        "residential": "No own-year FY2025-26 household survey; gas2025-07 through2026-03 (9months); liquid sales2025-07 through2026-06 (12months)",
        "waste": "No own-year management observations; disposal-history/population continuation is modelled, not observed monthly data",
        "ippu": "Physical QIM output2025-07 through2026-06 (12months); no FY2025-26 physical clinker/ammonia process activity established",
        "fugitives": "Gas consumption2025-07 through2026-03 (9months); no full-year production/infrastructure controls established",
        "railways": "Customer sales2025-07 through2026-06 (12months); no matching operator traction/non-traction consumption panel established",
        "aviation": "Resident-airline customer-class sales2025-07 through2026-06 (12months); no matching own-year domestic/international movements established",
    }
    rows = []
    for sector, (path, edition, usable, gap) in sources.items():
        for start in (2024, 2025):
            candidate = start == 2025
            rows.append({"sector": sector, "period": fiscal(start), "period_start": "%d-07-01" % start,
                "period_end": "%d-06-30" % (start + 1), "publication_edition": edition,
                "source_location": path, "source_location_status": "held",
                "variables_units": usable, "covered_months": "2024-07 through2025-06; annual source estimates or explicitly modelled continuation" if not candidate else candidate_months[sector],
                "revision_status": "current-vintage or explicit dated edition; numerical changes retained" if not candidate else "provisional; source boundaries retained",
                "actual_usability": usable if not candidate else gap,
                "release_class": "common_core" if not candidate else "separate_provisional_extension",
                "audit_date": "2026-10-14"})
    # Detailed source witnesses disambiguate full-year production, partial-year
    # sales, census snapshots and modelled continuation within the sector summary.
    witnesses = (
        ("transport;industry;aviation;railways", 2025, relative(OCAC / "Jul-Jun-2025-26-ENE-Pak-2.pdf"),
         "OCAC FY2025-26", "2025-07 through2026-06 (12months)", "t bycustomer/product", "usable provisional annual source panel; previous-year sector/product split unresolved"),
        ("transport;industry;aviation", 2025, relative(OCAC),
         "OCAC twelve direct monthly province editions", "2025-07 through2026-06 (12months)", "t bycustomer/product/province", "usable own-year allocation proxy; annual/monthly edition residuals retained"),
        ("ippu;industry", 2024, relative(QIM / "Table-3-2024-25-6.xlsx"),
         "PBS June2026 Table3 revision", "2024-07 through2025-06 (12months)", "physicalt;cementkt;glass1000m2", "usable core physical production; no ammonia or clinker row"),
        ("ippu;industry", 2025, relative(QIM / "Table-2-June-26.xlsx"),
         "PBS June2026 Table2 provisional", "2025-07 through2026-06 (12months)", "physicalt;cementkt;glass1000m2;fertilizernutrientt", "usable selected provisional processes; production not offtake/application"),
        ("agriculture", 2024, relative(AGRICULTURE),
         "Survey2025-26 Tables2.3/2.4", "FY2024-25 cropseason annualestimate", "ha;t;cotton1000bales", "usable revised core crops; historical overlapping revisions propagated"),
        ("agriculture", 2025, relative(AGRICULTURE),
         "Survey2025-26 Tables2.3/2.4", "FY2025-26 cropseason provisionalestimate", "ha;t;cotton1000bales", "usable provisional crop activity; not twelve observed monthly values"),
        ("agriculture", 2024, "data/nfdc/AFR_2024-25.pdf",
         "NFDC AFR2024-25 Annex5/28", "2024-07 through2025-06 (12months)", "fertilizerproductt;nutrientt;feed/fuelMMCF", "usable core monthly/annual product and nitrogen controls"),
        ("agriculture", 2025, relative(AGRICULTURE),
         "Survey2025-26 Table2.9", "2025-07 through2026-03 (9months)", "1000nutrientt", "partial only; cannot become full-year fertilizer application"),
        ("agriculture", 2025, "data/nfdc/FR_2026.pdf",
         "NFDC monthly review January2026", "2026-01 (1month); cumulativeOct2025-Jan2026", "product/nutrientt", "held, read; not an annual fertilizer review"),
        ("fugitives;residential;industry;transport", 2025, "data/raw/finance_division/survey_2025_26/14_Energy.pdf",
         "Survey2025-26 Table14.9", "2025-07 through2026-03 (9months)", "gas/RLNGMMCFD", "partial only; winter-weighted quantities not annualised"),
        ("industry", 2025, "data/raw/finance_division/survey_2025_26/14_Energy.pdf",
         "Survey2025-26 coalTable14.10", "2025-07 through2026-03 (9months)", "coal1000t", "partial consumer allocations; no full-year observed coal total"),
        ("agriculture", 2024, "data/shared/activity/pbs_census/livestock_province_2024.csv",
         "PBS AgriculturalCensus2024", "census stock snapshot, not monthlyactivity", "heads", "level anchored to the adopted census; year-to-year change is FAOSTAT relative stocks, which reproduce the Survey's 1996-2006 inter-census growth projection (Survey2025-26 Table2.5 footnote): no census-consistent growth bridge is held, and FY2024-25 continues the 2019-2024 projected growth; old Survey levels are not spliced"),
        ("residential;waste", 2024, "sources/HIES/HIES-2024-25-Report-Final-2.pdf",
         "HIES2024-25 heldmicrodata/questionnaire", "survey monthly recall; annualised surveyquantity, not12observed nationalmonths", "kg/household/month;surveyweights", "referenceS; report fieldwork dates differ across introductory sections; observation-period uncertainty retained"),
        ("power", 2024, relative(NEPRA),
         "NEPRA SoIR2025 Tables17/18/21/26", "2024-07 through2025-06 (12months)", "fuelmass;MMBTU;MMCFT;GWhnet", "coreCPPA/GENCO; KEfuel-specific alternatives; generation gaps explicitlyreconstructed"),
        ("residential", 2024, "data/raw/nasa_power/",
         "NASA POWER MERRA2 reanalysis monthlyAPI v2.10.0, five2024_2025 cityresponses", "2024-07 through2025-06 (12months)",
         "C;HDD18/CDD24 degree-days", "usable own-year seasonal proxy; monthly-mean reanalysis point temperature is not observed household fuel consumption"),
    )
    # rc68 audit (2026-10-15): documents read for the FY2025-26 candidacy, held and negative.
    audited = (
        ("agriculture", 2025, relative(NFDC_JULY_2026),
         "NFDC Monthly Fertilizer Review July2026 Annex-3", "2025-07 through2026-06 (12months)",
         "nutrient t by month", "usable provisional full-year nitrogen offtake; FY2024-25 column reproduces AFR2024-25; Jul-Mar reproduces Survey Table2.9; offtake, not application"),
        ("power", 2024, "data/raw/nepra/per_generation/nepra_per_generation_fy2024_25.pdf",
         "NEPRA Performance Evaluation Report of operational power plants FY2024-25", "2024-07 through2025-06 (12months)",
         "generation, capacity, EPP/CPP", "FY2024-25 cross-check only; no station fuel mass beyond SoIR2025"),
        ("power", 2025, "data/raw/nepra/listings/nepra_state_of_industry_reports_20261015.html",
         "NEPRA SoIR index read 2026-10-15", "no FY2025-26 edition published (latest SoIR2025 = FY2024-25)",
         "index", "absent: FY2025-26 station fuel cannot be established"),
        ("fugitives;industry;residential;transport", 2025, "data/raw/ogra/listings/ogra_annual_reports_20261015.html",
         "OGRA report index read 2026-10-15", "no 2025-26 Petroleum Industry Report or Annual Report published (latest 2024-25)",
         "index", "absent: full-year FY2025-26 gas deliveries/production/infrastructure not established"),
        ("fugitives", 2024, "data/raw/pbs/minerals_listing/pbs_energy_mining_20261015.html",
         "PBS energy-mining index read 2026-10-15", "no Mineral Production 2024-25 or 2025-26 file (latest 2023-24)",
         "index", "absent: FY2024-25 and FY2025-26 upstream timing stays T4 by source absence"),
    )
    for sector, start, path, edition, covered, variables, usable in audited:
        rows.append({"sector": sector, "period": fiscal(start), "period_start": "%d-07-01" % start,
            "period_end": "%d-06-30" % (start + 1), "publication_edition": edition,
            "source_location": path, "source_location_status": "held" if not path.endswith(".html") else "held_index",
            "variables_units": variables, "covered_months": covered,
            "revision_status": "edition and exact source coverage retained", "actual_usability": usable,
            "release_class": "common_core" if start == 2024 else "separate_provisional_extension",
            "audit_date": "2026-10-15"})
    for sector, start, path, edition, covered, variables, usable in witnesses:
        rows.append({"sector": sector, "period": fiscal(start), "period_start": "%d-07-01" % start,
            "period_end": "%d-06-30" % (start + 1), "publication_edition": edition,
            "source_location": path, "source_location_status": "held",
            "variables_units": variables, "covered_months": covered,
            "revision_status": "edition and exact source coverage retained", "actual_usability": usable,
            "release_class": "common_core" if start == 2024 else "separate_provisional_extension",
            "audit_date": "2026-10-14"})
    return rows


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args(argv)
    results = []
    def emit(path, rows, fields=None, banner=None):
        results.append(write_or_check(asset_path(ROOT, path), fields or list(rows[0]), rows,
                                      args.check, banner=banner))
    ag, candidate_crops = agriculture_rows()
    agpath = "data/activity/agriculture/es_fiscal_series.csv"
    old_ag = csv_at(agpath)
    revisions = []
    old_values = {(r["fiscal_end_year"], r["item"]): r for r in old_ag}
    # Persist rc67 differences on every regeneration, not merely the first run.
    baseline_path = ROOT / "data/shared/activity/extension/rc67_agricultural_revision_controls.csv"
    if baseline_path.exists():
        old_values = {(r["fiscal_end_year"], r["item"]): r for r in read_rows(baseline_path)}
    for row in ag:
        old = old_values.get((str(row["fiscal_end_year"]), row["item"]))
        if old and float(old["value"]) != row["value"]:
            revisions.append({"domain": "agriculture", "period": row["fiscal_label"], "item": row["item"],
                              "old": old["value"], "new": row["value"], "unit": row["unit"],
                              "reason": "Later official source revision", "source": row["source_table"]})
    emit(agpath, ag)
    anchors = csv_at("data/activity/crops_2024.csv")
    values = {r["item"]: r["value"] for r in ag if r["fiscal_end_year"] == 2024}
    for row in anchors:
        crop = row["crop"]
        row["area_ha"] = values[crop + "_area_ha"]
        row["production_t"] = values["cotton_bales_000"] * 485.9 if crop == "cotton" else values[crop + "_production_t"]
        row["source"] = "Pakistan Economic Survey 2025-26 appendix Tables2.3/2.4; " + relative(AGRICULTURE)
        row["note"] = "Revised FY2023-24 physical activity from current official vintage; cotton remains seed-mass conversion485.9t per1000bales."
    emit("data/activity/crops_2024.csv", anchors)
    qim_core = qim_records(QIM / "Table-3-2024-25-6.xlsx", 2024)
    qim_candidate = qim_records(QIM / "Table-2-June-26.xlsx", 2025)
    panel_path = "data/activity/monthly/pbs_qim_month_item.csv"
    existing = [r for r in csv_at(panel_path) if int(r["fiscal_year_start"]) < 2024]
    emit(panel_path, existing + qim_core, banner=qim_banner())
    national_path = "data/activity/ippu/national_activity.csv"
    national = [r for r in csv_at(national_path) if int(r["period_end_year"]) != 2025]
    for serial, product, unit, code in (("83", "Cement", "kt", "2.A.1"),
                                       ("85", "Billets and ingots", "t", "2.C.1"),
                                       ("58", "Soda ash", "t", "2.B.7")):
        item = next(r for r in qim_core if r["item_no"] == serial)
        national.append({"period_end_year": 2025, "period_basis": "fiscal", "ipcc_code": code,
                         "product": product, "quantity": item["annual_printed"], "unit": unit,
                         "source_id": "pbs_qim_fy2024_25", "source_reference": item["source_file"] + " item" + serial,
                         "coverage": "large_scale_manufacturing", "status": "estimate" if serial == "83" else "control",
                         "data_tier": "A", "note": "Physical production. Cement is not clinker; steel process route remains separately parameterised; soda route remains screened."})
    emit(national_path, national)
    ledger, fuel, generation, controls, bagasse_rows = power_products()
    emit("data/activity/power/plant_fuel_ledger.csv", ledger)
    emit("data/activity/power/fuel_consumed.csv", fuel)
    oldgen = [r for r in csv_at("data/activity/power/generation_by_fuel.csv") if int(r["fiscal_year"]) != 2024]
    emit("data/activity/power/generation_by_fuel.csv", oldgen + generation)
    old_bagasse = [r for r in csv_at("data/activity/power/bagasse_generation.csv") if int(r["fiscal_year"]) != 2024]
    emit("data/activity/power/bagasse_generation.csv", old_bagasse + bagasse_rows)
    ogra_path = RAW / "ogra/state_of_industry/ogra_petroleum_industry_report_2024_25.pdf"
    ogra_text = text(ogra_path)
    match = re.search(r"^\s*5\s+Captive Power\s+([\d.]+)\s+([\d.]+)\s+([\d.]+)\s+([\d.]+)\s+([\d.]+)", ogra_text, re.M)
    if match is None:
        raise ValueError("OGRA full-year captive sector row not found")
    sn, ss, independent, total, pct = map(float, match.groups())
    if abs(sn + ss + independent - total) > .001:
        raise ValueError("OGRA captive sector row fails utility sum")
    captive = [r for r in csv_at("data/activity/power/captive_gas.csv") if int(r["fiscal_year"]) != 2025]
    captive.append({"fiscal_year": 2025, "sngpl_mmcfd": sn, "ssgcl_mmcfd": ss,
                    "independent_mmcfd": independent, "total_mmcfd": total,
                    "pct_national_gas": pct, "status": "verified",
                    "source": "OGRA State of the Regulated Petroleum Industry2024-25 Table3.29; " + relative(ogra_path),
                    "note": "FullJul2024-Jun2025 captive electricity gas; separate from general industry and fertilizer feedstock. Not survey nine-month annualisation."})
    emit("data/activity/power/captive_gas.csv", captive)
    fertilizer_match = re.search(r"^\s*4\s+Fertilizer\s+([\d.]+)\s+([\d.]+)\s+([\d.]+)\s+([\d.]+)\s+([\d.]+)", ogra_text, re.M)
    if fertilizer_match is None:
        raise ValueError("OGRA full-year fertilizer sector row not found")
    fsn, fss, findependent, ftotal, _ = map(float, fertilizer_match.groups())
    if abs(fsn + fss + findependent - ftotal) > .001:
        raise ValueError("OGRA fertilizer sector row fails utility sum")
    fertilizer = [r for r in csv_at("data/activity/fugitives/fertilizer_gas_class.csv")
                  if int(r["fiscal_year"]) != 2025]
    fertilizer.append({"fiscal_year": 2025, "sngpl_mmcfd": fsn, "ssgcl_mmcfd": fss,
                       "independent_mmcfd": findependent, "total_mmcfd": ftotal,
                       "printed_label": "Fertilizer", "status": "verified",
                       "source": "OGRA State of the Regulated Petroleum Industry2024-25 Table3.29; " + relative(ogra_path),
                       "note": "FullJul2024-Jun2025 fertilizer gas class, feed and process fuel together; NFDC Annex28 supplies the observed feed/fuel split independently."})
    emit("data/activity/fugitives/fertilizer_gas_class.csv", fertilizer)
    aviation_spec = importlib.util.spec_from_file_location("extension_aviation_ingest", ROOT / "scripts/ingest_aviation.py")
    aviation_ingest = importlib.util.module_from_spec(aviation_spec)
    aviation_spec.loader.exec_module(aviation_ingest)
    aviation_movements, _ = aviation_ingest.source_rows()
    emit("data/activity/aviation/movements_airport.csv", aviation_movements)
    fuel_sales = [r for r in csv_at("data/activity/aviation/fuel_sales.csv") if int(r["fiscal_year"]) != 2024]
    aviation_source = next(r for r in csv_at("data/activity/aviation/ocac_annual_sources.csv")
                           if r["fiscal_year"] == "2024")
    from pkinv.sectors.aviation.annual import _extract_columns
    aviation_cells = _extract_columns(ROOT / "sources/Aviation/ocac_annual" / aviation_source["local_file"])
    for (customer, product), quantity in aviation_cells.items():
        fuel_sales.append({"fiscal_year": 2024, "customer_category": customer,
                           "fuel_type": product.replace(" ", ""), "tonnes": quantity,
                           "source": "OCAC Sectoral Sales of Energy Products FY2024-25 final public annual report",
                           "source_url": aviation_source["source_url"],
                           "source_note": "Published customer class, not an IPCC domestic/international flight boundary; own-year source edition retained separately from later aggregate revisions."})
    if len([r for r in fuel_sales if int(r["fiscal_year"]) == 2024]) != 15:
        raise ValueError("FY2024-25 aviation fuel cells incomplete")
    emit("data/activity/aviation/fuel_sales.csv", fuel_sales)
    emit("outputs/pakemis_power_extension_controls.csv", controls)
    emit("outputs/pakemis_extension_activity_revisions.csv", revisions,
         ["domain", "period", "item", "old", "new", "unit", "reason", "source"])
    # Candidate activities remain separated at the file boundary: no full-scope total.
    provisional = []
    for r in qim_candidate:
        if r["item_no"] not in ("83", "85", "58", "82", "1-3", "68-71"):
            continue
        provisional.append({"period": fiscal(2025), "month": r["month"], "area": "PAKISTAN", "sector": r["consuming_sector"] or "ippu",
            "activity": r["item"], "quantity": r["quantity"], "unit": r["unit"],
            "status": "provisional_physical_output", "source": r["source_file"],
            "boundary": "Product output; not automatically clinker/ammonia/combustion",
            "temporal_support": "own-year monthly production", "spatial_support": "national only"})
    for r in candidate_crops:
        provisional.append({"period": r["fiscal_label"], "month": "", "area": "PAKISTAN", "sector": "agriculture", "activity": r["item"],
            "quantity": r["value"], "unit": r["unit"], "status": r["basis"], "source": r["source_table"],
            "boundary": "Provisional full crop-season quantity; no full-year fertilizer application implied",
            "temporal_support": "annual crop season", "spatial_support": "national only"})
    nitrogen_rows, _nitrogen_total = nfdc_fy2026_nitrogen()
    provisional.extend(nitrogen_rows)
    ocac_rows, ocac_revisions, ocac_checks = ocac_extension()
    provisional.extend(ocac_rows)
    emit("outputs/pakemis_fy2025_26_provisional_activity.csv", provisional)
    emit("outputs/pakemis_ocac_extension_revisions.csv", ocac_revisions)
    emit("outputs/pakemis_extension_source_controls.csv", ocac_checks)
    emit("outputs/pakemis_extension_source_availability.csv", availability_rows())
    estimates = provisional_emissions(provisional)
    emit("outputs/pakemis_fy2025_26_provisional_emissions.csv", estimates)
    monthly_estimates = []
    for estimate in estimates:
        own = [r for r in provisional if r["month"] and r["area"] == "PAKISTAN"
               and r["activity"] == estimate["source_process"]]
        if not own and estimate["source_process"].startswith("Billets/Ingots"):
            own = [r for r in provisional if r["month"] and r["activity"] == "Billets/Ingots"]
        values = {r["month"]: float(r["quantity"]) for r in own}
        if set(values) != set(months(2025)):
            raise ValueError("Incomplete own-year monthly extension: %s" % estimate["source_process"])
        annual = sum(values.values())
        # A source with genuinely no annual sales emits zero in every month;
        # missing observations never enter this branch.
        for month, quantity in sorted(values.items()):
            row = dict(estimate, month=month, emissions_Gg=(estimate["emissions_Gg"] * quantity / annual if annual else 0.0),
                       temporal_support="T2: own-year sales/production; annual normalisation is explicit")
            monthly_estimates.append(row)
    emit("outputs/pakemis_fy2025_26_provisional_monthly.csv", monthly_estimates)
    print("Extension: %d fiscal agricultural cells; %d FY2025 physical month-item cells; %d candidate records" %
          (len(ag), len(qim_core), len(provisional)))
    return int(any(results)) if args.check else 0


if __name__ == "__main__":
    raise SystemExit(main())
