"""Independent flare and methane-plume observations; never silently annualized."""
from __future__ import annotations

from pkinv.assets import asset_path

from pkinv.activity.store import read_rows as _activity_rows
import csv
import hashlib
import io
import os
import zipfile
import math
import random
import re
from collections import defaultdict

from pkinv.paths import ROOT
from pkinv.core import documents
from pkinv.core.geography import PROVINCES
from pkinv.core.gis.grid import ADM1_PATH, ADM1_TO_PROVINCE
from pkinv.activity.store import read_rows

FLARE_SOURCE = asset_path(ROOT, "sources", "Fugitives",
                            "world_bank_flare_sites_2012_2025.xlsx")
MARS_SOURCE = asset_path(ROOT, "sources", "Fugitives", "unep_mars_plumes.zip")
# rc41: the engine reads the CONSUMED DERIVATIVES - the Pakistan rows of the two
# documents above, extracted with every column by scripts/ingest_fugitives_observations.py
# (data/consumed_derivatives.csv) - so a tree that registers the documents without
# holding them (the validation package) builds the same products. --check re-derives
# the tables wherever the documents are held.
ACTIVITY = asset_path(ROOT, "data", "activity", "fugitives")
FLARE_TABLE = asset_path(ACTIVITY, "flare_sites_pak.csv")
MARS_PLUMES_TABLE = asset_path(ACTIVITY, "mars_plumes_pak.csv")
MARS_SOURCES_TABLE = asset_path(ACTIVITY, "mars_sources_pak.csv")
MARS_TABLES = {"unep_methanedata_detected_plumes.csv": MARS_PLUMES_TABLE,
               "unep_methanedata_detected_sources.csv": MARS_SOURCES_TABLE}
FIELD_SUPPLY_SOURCE = asset_path(
    ROOT, "data", "activity", "fugitives", "field_gas_supply_2023_24.csv")

FIELD_ALIASES = {
    "qadir pur": "qadirpur", "qadirpur processed": "qadirpur",
    "qadirpur permeate for epql": "qadirpur", "qadirpur raw gas for lpl": "qadirpur",
    "zamzama i": "zamzama", "mari hrl": "mari", "mari goru b": "mari",
    "mari for engro fertilizer": "mari", "mari gas company limited": "mari",
    "mari petroleum ltd mari gas field": "mari", "ogdcl uch gas field": "uch",
    "ppl kandhkot": "kandhkot", "zarghun": "zarghun south",
    "tay": "tando allah yar", "rizq ewt": "rizq", "kalabagh 1a": "kalabagh",
    "manzalai kohat 2": "manzalai", "sofiya": "sofiya",
}


def sha256(path):
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def extract_flare_sites(path=FLARE_SOURCE):
    """The document's Pakistan rows, every column - the ingest script's extractor."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    header = [str(value) for value in next(sheet.iter_rows(values_only=True))]
    rows = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        if values[1] != "Pakistan":
            continue
        row = {header[index]: value for index, value in enumerate(values)}
        rows.append(row)
    return header, rows


def pakistan_flare_sites(path=None):
    return read_rows(path or FLARE_TABLE)


def flare_timeseries(path=None):
    sites = pakistan_flare_sites(path)
    result = []
    digest = documents.source_digest(FLARE_TABLE)
    for year in range(2012, 2026):
        active = [row for row in sites if float(row[str(year)] or 0.0) > 0]
        result.append({
            "year": year,
            "flare_volume_bcm": sum(float(row[str(year)] or 0.0) for row in sites),
            "active_flare_sites": len(active),
            "source_file": os.path.basename(FLARE_SOURCE),
            "source_sha256": digest,
            "source_role": "independent satellite-derived flaring observation",
        })
    return result


def flare_emission_memo(data, year=2024, path=None):
    volume_bcm = next(row["flare_volume_bcm"] for row in flare_timeseries(path)
                      if row["year"] == year)
    volume_e6m3 = volume_bcm * 1000.0
    gas_mass_Gg = volume_e6m3 * data.p("gas_density_kg_per_m3")
    methane_Gg = gas_mass_Gg * data.p("gas_ch4_fraction")
    efficiency = data.p("flare_combustion_efficiency")
    return {
        "year": year,
        "flare_volume_bcm": volume_bcm,
        "unburned_ch4_Gg": methane_Gg * (1.0 - efficiency),
        "combustion_co2_Gg": methane_Gg * efficiency * 44.0 / 16.0,
        "combustion_efficiency": efficiency,
        "interpretation": ("observation-based flare memo; not added to inclusive IPCC 2019 "
                           "production factors"),
    }


def extract_mars_rows(filename, path=MARS_SOURCE):
    """The archive member's Pakistan oil-and-gas rows, every column - the extractor."""
    with zipfile.ZipFile(path) as package:
        stream = io.TextIOWrapper(package.open(filename), encoding="utf-8-sig")
        reader = csv.DictReader(stream)
        rows = [row for row in reader
                if row["country"] == "Pakistan" and row["sector"] == "Oil and Gas"]
        return list(reader.fieldnames), rows


def _mars_rows(filename, path=None):
    return read_rows(path or MARS_TABLES[filename])


def pakistan_mars_plumes(path=None):
    return _mars_rows("unep_methanedata_detected_plumes.csv", path)


def pakistan_mars_sources(path=None):
    return _mars_rows("unep_methanedata_detected_sources.csv", path)


def mars_summary(path=None):
    plumes = pakistan_mars_plumes(path)
    sources = pakistan_mars_sources(path)
    fluxes = [float(row["ch4_fluxrate"]) for row in plumes if row["ch4_fluxrate"]]
    return {
        "oil_gas_sources": len(sources),
        "oil_gas_plumes": len(plumes),
        "persistent_sources": sum(row["persistency_category"].lower() == "persistent"
                                  for row in sources),
        "notified_sources": sum(row["notified"].lower() == "yes" for row in sources),
        "max_observed_flux_kg_h": max(fluxes) if fluxes else None,
        "source_file": os.path.basename(MARS_SOURCE),
        "source_sha256": documents.source_digest(MARS_PLUMES_TABLE),
        "interpretation": ("event/facility evaluation only; instantaneous detections are not "
                           "annualized into a national total"),
    }


def _province_lookup(points):
    """Attach public-boundary province names to point records."""
    import geopandas as gpd
    from shapely.geometry import Point

    boundaries = gpd.read_file(ADM1_PATH).to_crs("EPSG:4326")
    name_field = "shapeName" if "shapeName" in boundaries.columns else "name"
    aliases = dict(ADM1_TO_PROVINCE)
    aliases.update({"Islamabad Capital Territory": "Other", "Gilgit-Baltistan": "Other",
                    "Azad Kashmir": "Other"})
    output = []
    for source, lon, lat in points:
        point = Point(float(lon), float(lat))
        hit = boundaries[boundaries.geometry.covers(point)]
        province = "Unassigned" if hit.empty else aliases.get(
            str(hit.iloc[0][name_field]), str(hit.iloc[0][name_field]))
        output.append((source, province))
    return output


def geolocated_flare_sites(path=None):
    rows = [dict(row) for row in pakistan_flare_sites(path)]
    assignments = _province_lookup(
        [(index, row["Longitude"], row["Latitude"]) for index, row in enumerate(rows)])
    for index, province in assignments:
        rows[index]["province"] = province
    return rows


def geolocated_mars_sources(path=None):
    rows = [dict(row) for row in pakistan_mars_sources(path)]
    assignments = _province_lookup(
        [(index, row["lon"], row["lat"]) for index, row in enumerate(rows)])
    for index, province in assignments:
        rows[index]["province"] = province
    return rows


def _pearson(xs, ys):
    xbar, ybar = sum(xs) / len(xs), sum(ys) / len(ys)
    numerator = sum((x - xbar) * (y - ybar) for x, y in zip(xs, ys))
    denominator = math.sqrt(sum((x - xbar) ** 2 for x in xs) *
                            sum((y - ybar) ** 2 for y in ys))
    return numerator / denominator if denominator else float("nan")


def _ranks(values):
    order = sorted(range(len(values)), key=lambda index: values[index])
    ranks = [0.0] * len(values)
    for rank, index in enumerate(order, 1):
        ranks[index] = float(rank)
    return ranks


def provincial_evaluation(upstream_by_province, year=2024):
    """Evaluate upstream allocation against segment-relevant satellite geography."""
    provinces = [p for p in PROVINCES if p != "Other"]
    flares = geolocated_flare_sites()
    mars = geolocated_mars_sources()
    rows = []
    for province in provinces:
        p_flares = [row for row in flares if row["province"] == province]
        p_mars = [row for row in mars if row["province"] == province]
        rows.append({
            "province": province,
            "upstream_inventory_mt_co2e": float(upstream_by_province.get(province, 0.0)),
            "flare_volume_bcm": sum(float(row[str(year)] or 0.0) for row in p_flares),
            "active_flare_sites": sum(float(row[str(year)] or 0.0) > 0 for row in p_flares),
            "all_flare_sites": len(p_flares),
            "mars_sources": len(p_mars),
            "mars_plumes": sum(int(row["n_plumes_detected"]) for row in p_mars),
            "persistent_mars_sources": sum(
                row["persistency_category"].lower() == "persistent" for row in p_mars),
        })
    for field in ("upstream_inventory_mt_co2e", "flare_volume_bcm", "mars_sources",
                  "mars_plumes"):
        total = sum(row[field] for row in rows)
        for row in rows:
            row[field + "_share"] = row[field] / total if total else 0.0
    return rows


def provincial_statistics(rows):
    inventory = [row["upstream_inventory_mt_co2e_share"] for row in rows]
    output = []
    for field, label in (("flare_volume_bcm_share", "VIIRS 2024 flare-volume share"),
                         ("mars_sources_share", "MARS source-count share"),
                         ("mars_plumes_share", "MARS plume-count share")):
        observed = [row[field] for row in rows]
        inv_top = {rows[i]["province"] for i in sorted(
            range(len(rows)), key=lambda i: inventory[i], reverse=True)[:2]}
        obs_top = {rows[i]["province"] for i in sorted(
            range(len(rows)), key=lambda i: observed[i], reverse=True)[:2]}
        output.append({
            "comparison": label,
            "n_provinces": len(rows),
            "pearson_r": _pearson(inventory, observed),
            "spearman_r": _pearson(_ranks(inventory), _ranks(observed)),
            "top2_overlap_count": len(inv_top & obs_top),
            "top2_overlap_fraction": len(inv_top & obs_top) / 2.0,
            "inventory_top2": ";".join(sorted(inv_top)),
            "observation_top2": ";".join(sorted(obs_top)),
            "interpretation": ("upstream-geography evaluation only; four provinces give low "
                               "statistical power and do not test diffuse downstream leakage"),
        })
    return output


def _haversine_km(lon1, lat1, lon2, lat2):
    radius = 6371.0088
    phi1, phi2 = math.radians(lat1), math.radians(lat2)
    dphi = math.radians(lat2 - lat1)
    dlambda = math.radians(lon2 - lon1)
    a = math.sin(dphi / 2) ** 2 + math.cos(phi1) * math.cos(phi2) * math.sin(dlambda / 2) ** 2
    return 2 * radius * math.asin(math.sqrt(a))


def mars_flare_matches(year=2024):
    """Nearest-site diagnostic; spatial coincidence is not an annual flux comparison."""
    flares = geolocated_flare_sites()
    output = []
    for source in geolocated_mars_sources():
        ranked = sorted(
            ((_haversine_km(float(source["lon"]), float(source["lat"]),
                            float(site["Longitude"]), float(site["Latitude"])), site)
             for site in flares), key=lambda pair: pair[0])
        distance, site = ranked[0]
        output.append({
            "mars_source": source["source_name"],
            "mars_source_type": source["source_type"],
            "province": source["province"],
            "mars_longitude": float(source["lon"]),
            "mars_latitude": float(source["lat"]),
            "n_plumes_detected": int(source["n_plumes_detected"]),
            "nearest_flare_id": site["Flare id"],
            "nearest_field": site["Field name"],
            "nearest_operator": site["Operator"],
            "nearest_flare_longitude": float(site["Longitude"]),
            "nearest_flare_latitude": float(site["Latitude"]),
            "distance_km": distance,
            "within_10_km": distance <= 10.0,
            "within_25_km": distance <= 25.0,
            "nearest_flare_volume_bcm": float(site[str(year)] or 0.0),
        })
    return output


def _canonical_field(value):
    return _canonical_field_with_method(value)[0]


def _canonical_field_with_method(value):
    text = re.sub(r"[^a-z0-9]+", " ", str(value or "").lower()).strip()
    if not text or text in ("na", "0"):
        return text, "blank"
    if text in FIELD_ALIASES:
        return FIELD_ALIASES[text], "alias"
    return text, "exact_normalized"


def observation_match_audit(year=2024):
    """Expose every MARS-to-flare-to-field decision and its review status."""
    official_fields = {
        row["canonical_field"] for row in field_supply_rows()
        if row["fiscal_year"] == year
    }
    output = []
    for match in mars_flare_matches(year):
        canonical, method = _canonical_field_with_method(match["nearest_field"])
        operator = str(match.get("nearest_operator") or "").strip()
        source_type = str(match.get("mars_source_type") or "")
        distance = float(match["distance_km"])
        reasons = []
        if canonical not in official_fields:
            reasons.append("absent_official_field")
        if distance > 25.0:
            reasons.append("distance_gt_25_km")
        elif distance > 10.0:
            reasons.append("distance_gt_10_km")
        if "transmission" in source_type.lower():
            reasons.append("transmission_type_source")
        if not operator or operator.lower() in ("0", "na", "n/a", "unknown"):
            reasons.append("missing_operator")
        unmatched = ("absent_official_field" in reasons or
                     "distance_gt_25_km" in reasons)
        confidence = "unmatched" if unmatched else ("review_required" if reasons else "high")
        output.append({
            **match,
            "canonical_field": canonical,
            "canonicalization_method": method,
            "official_field_present": canonical in official_fields,
            "match_confidence": confidence,
            "review_reason_codes": ";".join(reasons),
            "audit_rule": ("high if official field, <=10 km, non-transmission and operator known; "
                           "review at 10-25 km/transmission/missing operator; unmatched if >25 km "
                           "or field absent"),
        })
    return output


def field_supply_rows(path=FIELD_SUPPLY_SOURCE):
    rows = list(_activity_rows(path, comments=False))
    for row in rows:
        row["fiscal_year"] = int(row["fiscal_year"])
        row["gas_supply_mmcfd"] = float(row["gas_supply_mmcfd"])
        row["canonical_field"] = _canonical_field(row["canonical_field"])
    return rows


def field_observation_evaluation(year=2024):
    """Compare official marketable field supply with named independent detections.

    This is a field-ranking evaluation, not a top-down national methane estimate.
    MARS events inherit the nearest World Bank flare field only where that flare has
    a usable name; no instantaneous plume flux is annualized.
    """
    supply = defaultdict(float)
    provinces = defaultdict(set)
    for row in field_supply_rows():
        if row["fiscal_year"] != year:
            continue
        supply[row["canonical_field"]] += row["gas_supply_mmcfd"]
        provinces[row["canonical_field"]].add(row["province"])

    flare_volume = defaultdict(float)
    flare_sites = defaultdict(int)
    for row in pakistan_flare_sites():
        field = _canonical_field(row["Field name"])
        if field in ("", "na", "0"):
            continue
        flare_volume[field] += float(row[str(year)] or 0.0)
        flare_sites[field] += 1

    mars_sources = defaultdict(int)
    mars_plumes = defaultdict(int)
    persistent = defaultdict(int)
    distances = defaultdict(list)
    for match in mars_flare_matches(year):
        field = _canonical_field(match["nearest_field"])
        if field in ("", "na", "0"):
            continue
        mars_sources[field] += 1
        mars_plumes[field] += int(match["n_plumes_detected"])
        distances[field].append(float(match["distance_km"]))
        source = next((row for row in geolocated_mars_sources()
                       if row["source_name"] == match["mars_source"]), None)
        if source and source["persistency_category"].lower() == "persistent":
            persistent[field] += 1

    fields = sorted(supply)
    total_supply = sum(supply.values())
    observed_supply = sum(supply[field] for field in fields
                          if flare_sites[field] or mars_sources[field])
    output = []
    for field in fields:
        output.append({
            "canonical_field": field,
            "province": ";".join(sorted(provinces[field])),
            "official_marketable_supply_mmcfd": supply[field],
            "official_supply_share": supply[field] / total_supply if total_supply else 0.0,
            "viirs_flare_volume_bcm": flare_volume[field],
            "viirs_named_sites": flare_sites[field],
            "mars_sources_assigned": mars_sources[field],
            "mars_plumes_assigned": mars_plumes[field],
            "persistent_mars_sources": persistent[field],
            "median_mars_to_flare_km": (_quantile(distances[field], 0.5)
                                         if distances[field] else ""),
            "has_independent_observation": bool(flare_sites[field] or mars_sources[field]),
            "official_supply_covered_by_named_observations_pct": (
                100.0 * observed_supply / total_supply if total_supply else 0.0),
            "evaluation_scope": ("field ranking of marketable gas supply versus named VIIRS/MARS "
                                 "evidence; diffuse downstream leakage is not tested"),
        })
    return output


def _quantile(values, probability):
    if not values:
        return float("nan")
    ordered = sorted(values)
    position = probability * (len(ordered) - 1)
    low = int(math.floor(position))
    high = int(math.ceil(position))
    if low == high:
        return ordered[low]
    return ordered[low] + (ordered[high] - ordered[low]) * (position - low)


def _average_ranks(values):
    order = sorted(range(len(values)), key=lambda index: values[index])
    ranks = [0.0] * len(values)
    start = 0
    while start < len(order):
        end = start + 1
        while end < len(order) and values[order[end]] == values[order[start]]:
            end += 1
        rank = (start + 1 + end) / 2.0
        for position in order[start:end]:
            ranks[position] = rank
        start = end
    return ranks


def _kendall_tau_b(xs, ys):
    concordant = discordant = ties_x = ties_y = 0
    for first in range(len(xs)):
        for second in range(first + 1, len(xs)):
            dx = (xs[first] > xs[second]) - (xs[first] < xs[second])
            dy = (ys[first] > ys[second]) - (ys[first] < ys[second])
            if dx == 0 and dy == 0:
                continue
            if dx == 0:
                ties_x += 1
            elif dy == 0:
                ties_y += 1
            elif dx == dy:
                concordant += 1
            else:
                discordant += 1
    denominator = math.sqrt((concordant + discordant + ties_x) *
                            (concordant + discordant + ties_y))
    return (concordant - discordant) / denominator if denominator else float("nan")


def field_observation_statistics(rows):
    supply = [float(row["official_marketable_supply_mmcfd"]) for row in rows]
    output = []
    for field, label in (("viirs_flare_volume_bcm", "VIIRS named-field flare volume"),
                         ("mars_sources_assigned", "MARS sources assigned by nearest flare"),
                         ("mars_plumes_assigned", "MARS plume count assigned by nearest flare")):
        observed = [float(row[field]) for row in rows]
        nonzero = sum(value > 0 for value in observed)
        top_n = min(10, len(rows))
        supply_top = {rows[index]["canonical_field"] for index in sorted(
            range(len(rows)), key=lambda index: supply[index], reverse=True)[:top_n]}
        observed_top = {rows[index]["canonical_field"] for index in sorted(
            range(len(rows)), key=lambda index: observed[index], reverse=True)[:top_n]
                       if observed[index] > 0}
        output.append({
            "comparison": label,
            "n_official_fields": len(rows),
            "n_fields_with_nonzero_observation": nonzero,
            "spearman_rho_with_ties": _pearson(_average_ranks(supply),
                                                 _average_ranks(observed)),
            "kendall_tau_b": _kendall_tau_b(supply, observed),
            "top10_overlap_count": len(supply_top & observed_top),
            "top10_observed_fields_available": len(observed_top),
            "supply_coverage_by_any_named_observation_pct": (
                float(rows[0]["official_supply_covered_by_named_observations_pct"])
                if rows else 0.0),
            "interpretation": ("field-level ranking diagnostic with zero-observation official "
                               "fields retained; observations are not annualized emissions"),
        })
    return output


def _top_overlap(supply, observed, top_n=10):
    top_n = min(top_n, len(supply))
    supply_top = set(sorted(range(len(supply)), key=lambda i: supply[i], reverse=True)[:top_n])
    observed_top = {
        index for index in sorted(
            range(len(observed)), key=lambda i: observed[i], reverse=True)[:top_n]
        if observed[index] > 0
    }
    return len(supply_top & observed_top), len(observed_top)


def field_observation_permutation(rows, n_iter=10000, seed=42):
    """Permutation tests for sparse field rankings, retaining all zero-observation fields."""
    supply = [float(row["official_marketable_supply_mmcfd"]) for row in rows]
    supply_ranks = _average_ranks(supply)
    province_groups = defaultdict(list)
    for index, row in enumerate(rows):
        province_groups[row["province"]].append(index)
    output = []
    comparisons = (
        ("viirs_flare_volume_bcm", "VIIRS named-field flare volume"),
        ("mars_sources_assigned", "MARS sources assigned by nearest flare"),
        ("mars_plumes_assigned", "MARS plume count assigned by nearest flare"),
    )
    for comparison_index, (field, label) in enumerate(comparisons):
        observed = [float(row[field]) for row in rows]
        observed_rho = _pearson(supply_ranks, _average_ranks(observed))
        observed_overlap, observed_available = _top_overlap(supply, observed)
        for null_index, null_model in enumerate(("national", "within_province")):
            run_seed = seed + 100 * comparison_index + null_index
            generator = random.Random(run_seed)
            rho_extreme = overlap_extreme = 0
            for _ in range(n_iter):
                if null_model == "national":
                    permuted = list(observed)
                    generator.shuffle(permuted)
                else:
                    permuted = list(observed)
                    for indices in province_groups.values():
                        values = [observed[index] for index in indices]
                        generator.shuffle(values)
                        for index, value in zip(indices, values):
                            permuted[index] = value
                rho = _pearson(supply_ranks, _average_ranks(permuted))
                overlap, _ = _top_overlap(supply, permuted)
                rho_extreme += abs(rho) >= abs(observed_rho)
                overlap_extreme += overlap >= observed_overlap
            output.append({
                "comparison": label,
                "null_model": null_model,
                "n_official_fields": len(rows),
                "n_fields_with_nonzero_observation": sum(value > 0 for value in observed),
                "observed_spearman_rho_with_ties": observed_rho,
                "spearman_two_sided_empirical_p": (rho_extreme + 1) / (n_iter + 1),
                "observed_top10_overlap_count": observed_overlap,
                "top10_observed_fields_available": observed_available,
                "top10_overlap_upper_tail_empirical_p": (
                    (overlap_extreme + 1) / (n_iter + 1)),
                "iterations": n_iter,
                "seed": run_seed,
                "interpretation": ("zeros retained; observation values permuted nationally or "
                                   "within province; no plume flux is annualized"),
            })
    return output
