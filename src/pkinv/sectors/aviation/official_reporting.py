"""Pakistan UNFCCC inventory entries, including explicit notation-key gaps."""
from __future__ import annotations

from pkinv.assets import asset_path

import csv
import hashlib
import io
import os
import zipfile

from pkinv.paths import DATA
from pkinv.paths import ROOT
from pkinv.core import documents
from pkinv.core.source_extracts import materialized
from pkinv.activity.store import read_rows

SOURCE = asset_path(DATA, "activity", "aviation", "unfccc_inventory_controls.csv")
CRT_SOURCE = asset_path(ROOT, "sources", "Aviation", "PAK_CRT_2024.zip")
CRT_WORKBOOK = "ENERGY-PAK-CRT.xlsx"


def _crt_payload(path=None):
    path = path or CRT_SOURCE
    documents.require(path, "the UNFCCC CRT tables are read here")
    with open(path, "rb") as handle:
        package_payload = handle.read()
    with zipfile.ZipFile(io.BytesIO(package_payload)) as package:
        workbook_payload = package.read(CRT_WORKBOOK)
    return path, package_payload, workbook_payload


def controls(path=None):
    rows = read_rows(path or SOURCE, comments=False)
    for row in rows:
        row["reported_inventory_year"] = int(row["reported_inventory_year"])
        row["matched_year"] = int(row["matched_year"])
        row["submission_year"] = int(row["submission_year"])
        row["value_co2e_t"] = (float(row["value_co2e_t"])
                                if row["value_co2e_t"].strip() else None)
        row["activity_tj"] = (float(row["activity_tj"])
                              if row["activity_tj"].strip() else None)
    return rows


def comparison(controls_, oecd_rows):
    oecd = {}
    for row in oecd_rows:
        calendar_year = int(row["month"][:4])
        fiscal_year = int(row["fiscal_year"])
        for basis, year in (("calendar", calendar_year), ("fiscal", fiscal_year)):
            key = (basis, year, "international bunker")
            oecd[key] = oecd.get(key, 0.0) + float(row["territory_international_co2_t"])
            key = (basis, year, "domestic territory")
            oecd[key] = oecd.get(key, 0.0) + float(row["territory_domestic_co2_t"])
    output = []
    for row in controls_:
        value = row["value_co2e_t"]
        benchmark = oecd.get((row["period_basis"], row["matched_year"], row["scope"]))
        output.append({
            **row,
            "oecd_matched_period_co2_t": benchmark if benchmark is not None else "",
            "official_over_oecd_ratio": (value / benchmark
                                          if value is not None and benchmark else ""),
            "comparison_interpretation": (
                "official national reporting comparison" if value is not None
                else "notation-key gap; no inferred value"),
        })
    return output


@materialized("aviation_extract_crt")
def extract_crt(path=None):
    """Extract aviation values from the official Pakistan 2022 CRT workbook."""
    from openpyxl import load_workbook

    path, payload, workbook_payload = _crt_payload(path)
    workbook = load_workbook(io.BytesIO(workbook_payload), data_only=True,
                             read_only=True)
    summary = workbook["Table1"]
    values = {}
    for row in summary.iter_rows(values_only=True):
        category = str(row[1] or "").strip()
        if category in ("1.A.3.a. Domestic aviation", "1.D.1.a. Aviation"):
            values[category] = row
    domestic_activity = None
    for row in workbook["Table1.A(a)s3"].iter_rows(values_only=True):
        if str(row[1] or "").strip().startswith("1.A.3.a.  Domestic aviation"):
            domestic_activity = float(row[2])
    international_activity = None
    distribution = {}
    for row in workbook["Table1.D"].iter_rows(values_only=True):
        label = str(row[1] or "").strip()
        if label.startswith("1.D.1.a. International aviation"):
            international_activity = float(row[2])
        if len(row) > 11 and row[10] == "Aviation":
            distribution = {"domestic_pct": float(row[11]),
                            "international_pct": float(row[12])}
    if len(values) != 2 or domestic_activity is None or international_activity is None:
        raise ValueError("official CRT aviation rows are incomplete")
    output = []
    for scope, category, activity in (
            ("domestic territory", "1.A.3.a. Domestic aviation", domestic_activity),
            ("international bunker", "1.D.1.a. Aviation", international_activity)):
        row = values[category]
        jet_notation = (workbook["Table1.A(a)s3"]["C18"].value if scope == "domestic territory"
                        else workbook["Table1.D"]["C11"].value)
        output.append({
            "reported_inventory_year": int(summary["J1"].value),
            "scope": scope, "category": category,
            "co2_t": float(row[2]) * 1000.0,
            "ch4_t": float(row[3]) * 1000.0,
            "n2o_t": float(row[4]) * 1000.0,
            "total_co2e_t": float(row[9]) * 1000.0,
            "activity_tj": activity, "fuel_reported": "Aviation gasoline",
            "jet_kerosene_notation": jet_notation,
            "domestic_distribution_pct": distribution["domestic_pct"],
            "international_distribution_pct": distribution["international_pct"],
            "source_file": os.path.basename(path),
            "source_sha256": hashlib.sha256(payload).hexdigest(),
            "workbook_file": CRT_WORKBOOK,
            "workbook_sha256": hashlib.sha256(workbook_payload).hexdigest(),
        })
    return output


@materialized("aviation_extract_crt_audit")
def extract_crt_audit(path=None):
    """Cell-level provenance for every aviation value interpreted from the CRT."""
    from openpyxl import load_workbook

    path, package_payload, workbook_payload = _crt_payload(path)
    workbook = load_workbook(io.BytesIO(workbook_payload), data_only=False,
                             read_only=False)
    specifications = [
        ("Table1", "C25", "domestic_co2", "kt CO2", "multiply by 1000 to t CO2"),
        ("Table1", "D25", "domestic_ch4", "kt CH4", "multiply by 1000 to t CH4"),
        ("Table1", "E25", "domestic_n2o", "kt N2O", "multiply by 1000 to t N2O"),
        ("Table1", "J25", "domestic_total_co2e", "kt CO2e", "multiply by 1000 to t CO2e"),
        ("Table1", "C53", "international_co2", "kt CO2", "multiply by 1000 to t CO2"),
        ("Table1", "D53", "international_ch4", "kt CH4", "multiply by 1000 to t CH4"),
        ("Table1", "E53", "international_n2o", "kt N2O", "multiply by 1000 to t N2O"),
        ("Table1", "J53", "international_total_co2e", "kt CO2e", "multiply by 1000 to t CO2e"),
        ("Table1.A(a)s3", "C16", "domestic_activity_total", "TJ", "retain as TJ"),
        ("Table1.A(a)s3", "C17", "domestic_aviation_gasoline_activity", "TJ", "retain as TJ"),
        ("Table1.A(a)s3", "H17", "domestic_aviation_gasoline_co2", "kt CO2", "multiply by 1000 to t CO2"),
        ("Table1.A(a)s3", "C18", "domestic_jet_kerosene_activity", "notation key", "retain notation key"),
        ("Table1.A(a)s3", "H18", "domestic_jet_kerosene_co2", "notation key", "retain notation key"),
        ("Table1.D", "C10", "international_activity_total", "TJ", "retain as TJ"),
        ("Table1.D", "G10", "international_co2_total", "kt CO2", "multiply by 1000 to t CO2"),
        ("Table1.D", "C11", "international_jet_kerosene_activity", "notation key", "retain notation key"),
        ("Table1.D", "G11", "international_jet_kerosene_co2", "notation key", "retain notation key"),
        ("Table1.D", "C12", "international_aviation_gasoline_activity", "TJ", "retain as TJ"),
        ("Table1.D", "G12", "international_aviation_gasoline_co2", "kt CO2", "multiply by 1000 to t CO2"),
        ("Table1.D", "L9", "domestic_distribution", "percent", "retain exact percentage; round only for display"),
        ("Table1.D", "M9", "international_distribution", "percent", "retain exact percentage; round only for display"),
    ]
    rows = []
    for sheet_name, coordinate, field, unit, transformation in specifications:
        cell = workbook[sheet_name][coordinate]
        value = cell.value
        transformed = value
        if isinstance(value, (int, float)) and unit.startswith("kt "):
            transformed = float(value) * 1000.0
        rows.append({
            "source_zip_file": os.path.basename(path),
            "source_zip_sha256": hashlib.sha256(package_payload).hexdigest(),
            "workbook_file": CRT_WORKBOOK,
            "workbook_sha256": hashlib.sha256(workbook_payload).hexdigest(),
            "worksheet": sheet_name, "cell": coordinate,
            "source_reference": "%s!%s" % (sheet_name, coordinate),
            "field": field, "raw_value": value, "raw_data_type": cell.data_type,
            "unit": unit, "transformation": transformation,
            "transformed_value": transformed,
            "formula": value if cell.data_type == "f" else "",
            "notation_key": value if isinstance(value, str) and
            ("NE" in value or "NO" in value or "NA" in value) else "",
        })
    return rows
